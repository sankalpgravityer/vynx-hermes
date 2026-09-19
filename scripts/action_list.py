#!/usr/bin/env python
"""The products worth acting on, and what each one would cost to fix.

    python scripts/action_list.py reports/generated/dossier-BEFORE-*.json \
        --out reports/generated/midtex-action-list.xlsx

WHAT IT IS FOR. `audit_approved_window.py` answers "what is wrong with this
shelf" across every product, which is the right question once. Having answered
it, the next question is narrower and is the one somebody actually works from:
*which* products have a picture problem or a blocking rule, what would the chain
do to each, and what does that cost. A 408-row sheet where 235 rows say
"MANUAL_APPROVAL, DRIFT.001" does not answer that; this does.

WHY IT READS THE DOSSIER JSON rather than the database. The dossier already ran
the rules, the cut-out measurement, the image gate and the photo audit, and its
verdicts are the FULL-STRENGTH ones — `check_product.inspect` passes the
product's size and the gender its master category implies, which is what the
chain's own gate step passes. Re-deriving any of that here would be a second
implementation of the same judgement and a second chance to get it subtly wrong,
which is exactly what happened to the audit script's gate column on 18 Sep 2026.
So: no database, no vision calls, nothing to spend. The JSONs are the input.

THE THREE SIGNALS, kept apart on purpose. A product can have any combination,
and they are repaired by different steps at different prices:

    gate          the lead render is refused. Re-renders; the whole set when the
                  model's gender or build is wrong, since one model carries them all.
    photo audit   a render or a cut-out is defective. Re-renders the named views,
                  or re-cuts the named cut-outs.
    rules         a blocking finding in the data. Mostly free to repair.

Column `Signals` says which of the three fired, so the sheet can be filtered to
"all three" or to any one of them.
"""
from __future__ import annotations

import argparse
import glob
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# The renderer's own order, and the fallback the gate uses when policy is silent.
ALL_VIEWS = ["AI_FRONT_34", "AI_BACK_34", "AI_FRONT", "AI_BACK", "AI_CLOSEUP"]

# A refusal that is a property of the MODEL, not of one picture: one model wears
# every view, so redoing a single one would put a different person in it. Mirrors
# app/imaging/quality_gate.regen_views.
WHOLE_SET_CODES = ("MODEL_GENDER_MISMATCH", "BODY_SIZE_MISMATCH")

HDR = "1F3864"
FILL = {"high": "FDECEA", "medium": "FFF8E1", "low": "E8F5E9", "none": "FFFFFF"}


def paid_views(p: dict[str, Any], all_views: list[str]) -> list[str]:
    """Which views the chain would pay to render again for this product.

    The union of what the gate asks for and what the photo audit asks for — a
    product refused for gender AND carrying a close-up defect renders five views,
    not six. `review` verdicts ask for nothing: the gate returns those for a
    person to decide (a category that disagrees with the picture is usually the
    category being wrong, not the render), and re-rendering would spend money on
    a question nobody has answered yet.
    """
    wanted: set[str] = set()

    gate = p.get("judged") or {}
    if gate.get("action") == "regen":
        if gate.get("code") in WHOLE_SET_CODES:
            wanted.update(all_views)
        else:
            wanted.update(gate.get("bad_views") or [])
            # The lead is dragged in only by codes that are about the lead. The
            # two that name OTHER pictures never pull it along.
            if gate.get("lead_view") and gate.get("code") not in ("IMAGE_COMPOSITION", "RENDER_DEFECT"):
                wanted.add(str(gate["lead_view"]))

    photos = p.get("audited") or {}
    if photos.get("action") == "regen":
        wanted.update(photos.get("bad_views") or [])

    return [v for v in all_views if v in wanted] + sorted(v for v in wanted if v not in all_views)


def verdict_text(v: dict[str, Any] | None) -> str:
    if not v:
        return "not judged"
    if v.get("unavailable"):
        return f'COULD NOT RUN — {"; ".join(v.get("reasons") or [])[:120]}'
    head = f'{v.get("action")} {v.get("code") or ""}'.strip()
    why = "; ".join(v.get("reasons") or [])
    return (head + (f' — {why}' if why else ""))[:300]


def row_for(p: dict[str, Any], price: float, all_views: list[str]) -> dict[str, Any] | None:
    """One product's row, or None when it needs nothing."""
    gate = p.get("judged") or {}
    photos = p.get("audited") or {}
    cut = p.get("measured") or {}
    issues = p.get("issues") or []

    gate_bad = gate.get("action") in ("regen", "review") and not gate.get("unavailable")
    # A cut-out defect reaches us through the photo audit's soft flags as well as
    # its action, which is why `bad_cutouts` is checked and not only `action`.
    photos_bad = (photos.get("action") in ("regen", "review")
                  or bool(photos.get("bad_cutouts")))
    blocking = [i for i in issues
                if i.get("blocking") and not str(i["rule_id"]).startswith(("GATE:", "PHOTOS:", "CUTOUT:"))]

    if not (gate_bad or photos_bad or blocking):
        return None

    signals = ",".join(s for s, on in (("gate", gate_bad), ("photos", photos_bad),
                                       ("rules", bool(blocking))) if on)
    views = paid_views(p, all_views)
    worst = "none"
    for i in issues:
        if i["severity"] == "high" or (worst != "high" and i["severity"] == "medium"):
            worst = i["severity"] if i["severity"] == "high" else "medium"

    return {
        "Signals": signals,
        "All three": "yes" if signals.count(",") == 2 else "",
        "Worst": worst,
        "SKU": p.get("sku"),
        "Title": (p.get("title") or "")[:80],
        "Image gate": verdict_text(gate) if gate_bad else "ok",
        "Photo audit": verdict_text(photos) if photos_bad else "ok",
        "Bad cut-outs": ", ".join(photos.get("bad_cutouts") or cut.get("bad_views") or []) or "—",
        "Blocking rules": ", ".join(sorted({str(i["rule_id"]) for i in blocking})) or "—",
        "Blocking detail": " | ".join(f'{i["rule_id"]}: {i["message"]}' for i in blocking)[:600] or "—",
        "Views to re-render": ", ".join(views) or "—",
        "Paid views": len(views),
        "Render cost": round(len(views) * price, 3),
        "Product id": p.get("product_id"),
        "Edit URL": p.get("edit_url") or "",
    }


COLUMNS = [("Signals", 16), ("All three", 10), ("Worst", 9), ("SKU", 14), ("Title", 44),
           ("Image gate", 54), ("Photo audit", 54), ("Bad cut-outs", 16),
           ("Blocking rules", 30), ("Blocking detail", 70), ("Views to re-render", 30),
           ("Paid views", 11), ("Render cost", 12), ("Product id", 38), ("Edit URL", 42)]


def build(paths: list[Path], out: Path, price: float) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.config import policy

    all_views = list((policy().get("imagery") or {}).get("all_views") or ALL_VIEWS)

    products: list[dict[str, Any]] = []
    for f in paths:
        products += json.loads(f.read_text(encoding="utf-8"))["products"]

    rows = [r for r in (row_for(p, price, all_views) for p in products) if r]
    rows.sort(key=lambda r: ({"high": 0, "medium": 1, "low": 2, "none": 3}[r["Worst"]],
                             -r["Paid views"], str(r["SKU"])))

    wb = Workbook()
    hf, hfont = PatternFill("solid", fgColor=HDR), Font(bold=True, color="FFFFFF")

    # ---- Summary ----------------------------------------------------------- #
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Products needing action").font = Font(bold=True, size=14)
    sig = Counter(s for r in rows for s in r["Signals"].split(","))
    views = sum(r["Paid views"] for r in rows)
    meta = [
        ("Source", ", ".join(f.name for f in paths)),
        ("Products examined", len(products)),
        ("Products needing action", len(rows)),
        ("  image gate refuses", sig.get("gate", 0)),
        ("  photo audit finds a defect", sig.get("photos", 0)),
        ("  a data rule blocks", sig.get("rules", 0)),
        ("  all three at once", sum(1 for r in rows if r["All three"])),
        ("Paid render views", views),
        ("Render cost", f"${views * price:.2f}  (batch API 50% off: ${views * price / 2:.2f})"),
        ("Price assumed", f"${price:.3f} per view"),
        ("Method", "read from the dossier JSON — no database, no vision calls, nothing spent"),
    ]
    for n, (k, v) in enumerate(meta, start=3):
        ws.cell(row=n, column=1, value=k).font = Font(bold=True)
        ws.cell(row=n, column=2, value=v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95

    r0 = 3 + len(meta) + 1
    for i, name in enumerate(("Blocking rule", "Products"), start=1):
        c = ws.cell(row=r0, column=i, value=name)
        c.fill, c.font = hf, hfont
    counts = Counter(rid for r in rows for rid in r["Blocking rules"].split(", ") if rid != "—")
    for n, (rid, cnt) in enumerate(counts.most_common(), start=r0 + 1):
        ws.cell(row=n, column=1, value=rid)
        ws.cell(row=n, column=2, value=cnt)

    # ---- Action list -------------------------------------------------------- #
    ws = wb.create_sheet("Action list")
    for i, (name, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = hf, hfont
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    for n, r in enumerate(rows, start=2):
        for col, (name, _) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=n, column=col, value=r[name])
            c.alignment = Alignment(vertical="top", wrap_text=name in
                                    ("Title", "Image gate", "Photo audit", "Blocking detail"))
        ws.cell(row=n, column=3).fill = PatternFill("solid", fgColor=FILL[r["Worst"]])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 2)}"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json", nargs="+", help="dossier JSON file(s); globs are expanded")
    ap.add_argument("--out", required=True)
    ap.add_argument("--price", type=float, default=0.067,
                    help="price of one rendered view (default 0.067, gemini-3.1-flash-image at 1K)")
    args = ap.parse_args(argv)

    paths = [Path(p) for pat in args.json for p in sorted(glob.glob(pat))] or [Path(p) for p in args.json]
    missing = [p for p in paths if not p.exists()]
    if missing:
        sys.exit(f"not found: {', '.join(str(m) for m in missing)}")

    out = build(paths, Path(args.out), args.price)
    print(f"{out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
