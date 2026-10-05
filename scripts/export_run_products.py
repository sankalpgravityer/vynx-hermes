#!/usr/bin/env python
"""Export one auto-approval run's products to Excel: which were verified, held for a
human, or failed, and why. READ ONLY.

    python scripts/export_run_products.py --run adf6048a --prod
    python scripts/export_run_products.py --run adf6048a-15a0-4981-ad8e-9e0fafd3e38d --db <dsn>

`--run` takes the run id or a unique prefix of it. `--prod` reads the vnyx-prod
connection string from E:\\vnyx\\vnyx-api\\.env (or $VNYX_PROD_DATABASE_URL); without
it, Hermes' .env DATABASE_URL is used. Either way the connection is opened read-only
and the string is never printed.

Writes reports/<tenant>_run_<local start>_<run id prefix>.xlsx with five sheets:
Summary, All products, Verified, Held for human, Failed. Times are shown in --tz
(default Asia/Kolkata, the zone the Runs screen was read in).
"""
from __future__ import annotations

import argparse
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
API_ENV = Path(r"E:\vnyx\vnyx-api\.env")
UI = "https://try.vnyx.ai"

OUTCOME = {"VERIFIED": "Verified", "HELD_FOR_HUMAN": "Held for human", "FAILED": "Failed",
           "CANCELLED": "Withdrawn", "QUEUED": "In the queue", "RUNNING": "Running"}
ORDER = ["Verified", "Held for human", "Failed", "Withdrawn", "In the queue", "Running"]
FILL = {"Verified": "E2F0D9", "Held for human": "FFF2CC", "Failed": "F8CBAD"}
SCOPE = {"MANUAL_SINGLE": "Single product", "MANUAL_ID_LIST": "ID list",
         "MANUAL_DATE_RANGE": "By date", "MANUAL_FULL_REVIEW": "Whole Review section",
         "MANUAL_TAB": "Whole section", "MANUAL_RETRY": "Retry", "SCHEDULED": "Scheduled",
         "ON_ARRIVAL": "On arrival"}

# What the rule ids and preflight problems this run can carry mean, in words
# (app/rules/gate.py, vnyx-api scripts/approve-products.ts).
MEANING = {
    "PRICE.001": "the price is outside the grade's price window",
    "TAX.004": "the gender disagrees with the master category",
    "TAX.005": "the mannequin is wrong for the category",
    "SIZE.011": "the sizing guide is for the other gender",
    "DATA.010": "a required field is empty",
    "SIZE.001": "the size fields disagree with each other",
    "SIZE.002": "the EU size is not the one the sizing guide pairs with the size",
    "TEXT.003": "the title advertises a different size than the size field",
    "SIZE.013": "the sizing guide does not list this product's size",
    "SIZE.014": "the sizing guide is for the wrong half of the body",
    "SIZE.010": "no size chart covers this product",
    "DRIFT.001": "a column and its property copy disagree (category, subcategory or size)",
    "GENDER.001": "gender is undecided",
    "IMG.030": "no care-label photo",
    "IMG.040": "the photo shows a different kind of garment than the category",
    "no EU size": "EU size is empty (required to list)",
}

SQL = '''
SELECT rp.status::text, rp."productSku", rp."productTitle", p.title, rp.reason,
       rp."blockingRules", rp."preflightProblems", rp.approved,
       b.name, p."internationalSize", coalesce(p.properties->>'eu_size', p.properties->>'euSize'),
       p."sizingGuide", p."masterCategory", p.category, p."subCategory",
       p."currentStage"::text, p.status::text, p."shopifyProductId",
       rp.steps, rp."startedAt", rp."completedAt", rp."durationMs", rp.attempts, rp.error,
       rp."productId"::text
  FROM "AutoApprovalRunProduct" rp
  LEFT JOIN "Product" p ON p.id = rp."productId"
  LEFT JOIN "Brand" b ON b.id = p."brandId"
 WHERE rp."runId" = %s
'''

COLUMNS = [  # header, width
    ("Outcome", 15), ("SKU", 13), ("Title at run start", 44), ("Title now", 44),
    ("Why", 60), ("Rules still open", 26), ("Preflight problems", 18), ("Moved to Approved", 11),
    ("Brand now", 16), ("Size now", 9), ("EU size now", 9), ("Sizing guide now", 22),
    ("Master category", 11), ("Category", 11), ("Subcategory now", 15),
    ("Stage now", 11), ("Status now", 10), ("On Shopify", 9),
    ("What the run did", 70), ("Started", 16), ("Finished", 16), ("Took (min)", 9),
    ("Attempts", 8), ("Error", 30), ("Product id", 38), ("Edit link", 30),
]


def prod_dsn() -> str:
    if os.environ.get("VNYX_PROD_DATABASE_URL"):
        return os.environ["VNYX_PROD_DATABASE_URL"]
    for line in API_ENV.read_text(encoding="utf-8").splitlines():
        m = re.match(r'^\s*#?\s*DATABASE_URL\s*=\s*"?([^"\s]+/vnyx-prod(?:\?[^"\s]*)?)"?', line)
        if m:
            return m.group(1)
    raise SystemExit("no prod connection string: set VNYX_PROD_DATABASE_URL or pass --db")


def env_dsn() -> str:
    m = re.search(r'^\s*DATABASE_URL\s*=\s*"?([^"\s]+)', (ROOT / ".env").read_text(encoding="utf-8"), re.M)
    if not m:
        raise SystemExit("no DATABASE_URL in Hermes' .env — pass --db")
    return m.group(1)


def dedupe(items) -> list[str]:
    return list(dict.fromkeys(str(i) for i in (items or [])))


def why(outcome: str, reason: str | None, rules: list[str], preflight: list[str]) -> str:
    """The reason in words. Verified and failed rows carry a sentence already; a held
    row carries rule ids, which are spelled out."""
    if outcome == "Held for human":
        parts = [f"{r}: {MEANING.get(r, 'see the rule')}" for r in dedupe(rules)]
        parts += [MEANING.get(p, p) for p in dedupe(preflight)]
        return "; ".join(parts) or (reason or "")
    return reason or ""


def did(steps) -> str:
    """One line per step that ran: what it found or changed."""
    lines = []
    for s in steps or []:
        if not s.get("ran"):
            continue
        note = re.sub(r"\s+", " ", str(s.get("note") or "")).strip()
        if len(note) > 160:
            note = note[:157] + "..."
        flag = "" if s.get("ok", True) else " [FAILED]"
        lines.append(f"{s.get('step')}{flag}: {note}")
    return "\n".join(lines)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True, help="run id, or a unique prefix of it")
    ap.add_argument("--prod", action="store_true", help="read the vnyx-prod database (read-only)")
    ap.add_argument("--db", help="connection string (overrides --prod and Hermes' .env)")
    ap.add_argument("--tz", default="Asia/Kolkata", help="zone the times are shown in")
    ap.add_argument("--out", type=Path, default=ROOT / "reports")
    args = ap.parse_args()

    tz = ZoneInfo(args.tz)
    read_at = datetime.now(timezone.utc)
    dsn = args.db or (prod_dsn() if args.prod else env_dsn())
    with psycopg.connect(dsn, connect_timeout=30,
                         options="-c default_transaction_read_only=on -c statement_timeout=120000") as c:
        c.read_only = True
        db = c.execute("SELECT current_database()").fetchone()[0]
        runs = c.execute(
            '''SELECT r.id::text, t.name, t.id::text, r.source::text, r.status::text, r."startedAt",
                      r."completedAt", r."totalProducts", r."verifiedCount", r."heldCount",
                      r."failedCount", r."cancelledCount", r."approvedCount", r."completionReason",
                      r."configSnapshot"->>'mode'
                 FROM "AutoApprovalRun" r JOIN "Tenant" t ON t.id = r."tenantId"
                WHERE r.id::text LIKE %s''', (args.run.lower() + "%",)).fetchall()
        if len(runs) != 1:
            raise SystemExit(f"run {args.run!r} matched {len(runs)} runs")
        (run_id, tenant, tenant_id, source, status, started, completed, total, n_ver, n_held,
         n_fail, n_canc, n_appr, reason, mode) = runs[0]
        rows = c.execute(SQL, (run_id,)).fetchall()

    def local(ts):
        return ts.replace(tzinfo=timezone.utc).astimezone(tz).replace(tzinfo=None) if ts else None

    table = []
    for r in rows:
        (st, sku, title0, title, reason_txt, rules, preflight, approved, brand, size, eu, guide,
         master, cat, sub, stage, pstatus, shop, steps, s0, s1, ms, attempts, err, pid) = r
        outcome = OUTCOME.get(st, st)
        table.append([
            outcome, sku, title0, title, why(outcome, reason_txt, rules, preflight),
            ", ".join(dedupe(rules)), ", ".join(dedupe(preflight)), "Yes" if approved else "No",
            brand, size, eu, guide, master, cat, sub, stage, pstatus, "Yes" if shop else "No",
            did(steps), local(s0), local(s1), round(ms / 60000, 1) if ms else None,
            attempts, (err or "")[:500] or None, pid,
            f"{UI}/product/{pid}/edit?tenantId={tenant_id}" if pid else None,
        ])
    rank = {o: i for i, o in enumerate(ORDER)}
    table.sort(key=lambda t: (rank.get(t[0], 99), t[1] or ""))

    font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F4E5A")
    link = Font(name="Arial", size=10, color="0563C1", underline="single")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    headers = [h for h, _ in COLUMNS]
    zone = args.tz.split("/")[-1]
    shown = [f"{h} ({zone})" if h in ("Started", "Finished") else h for h in headers]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    def sheet(name: str, data: list[list]) -> None:
        ws = wb.create_sheet(name)
        ws.append(shown)
        for cell in ws[1]:
            cell.font, cell.fill = head, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for i, row in enumerate(data, start=2):
            ws.append(row)
            for j, cell in enumerate(ws[i], start=1):
                cell.font = font
                cell.alignment = wrap if headers[j - 1] in ("Why", "What the run did", "Error",
                                                            "Title at run start", "Title now") else top
            oc = ws.cell(i, 1)
            if oc.value in FILL:
                oc.fill = PatternFill("solid", start_color=FILL[oc.value])
                oc.font = bold_font
            for h in ("Started", "Finished"):
                ws.cell(i, headers.index(h) + 1).number_format = "yyyy-mm-dd hh:mm"
            lc = ws.cell(i, headers.index("Edit link") + 1)
            if lc.value:
                lc.hyperlink, lc.font = lc.value, link
        for j, (_, w) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(data) + 1)}"

    sheet("All products", table)
    for o in ("Verified", "Held for human", "Failed"):
        sheet(o, [t for t in table if t[0] == o])

    # --- Summary: counts are formulas over "All products"; the run's own counters
    # sit beside them as the cross-check.
    last = len(table) + 1
    rng = lambda h: f"'All products'!${get_column_letter(headers.index(h) + 1)}$2:" \
                    f"${get_column_letter(headers.index(h) + 1)}${max(2, last)}"
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 22
    lines: list[tuple] = [
        ("Run", None, None),
        ("Run id", run_id, None),
        ("Tenant", tenant, None),
        ("Scope", SCOPE.get(source, source), None),
        ("Mode", mode or "", None),
        (f"Started ({zone})", local(started), None),
        (f"Finished ({zone})", local(completed), None),
        ("Run status", f"{status} — {reason or ''}", None),
        (None, None, None),
        ("Products", "From this file", "Run counters (vnyx)"),
        ("Checked", f"=COUNTA({rng('SKU')})", total),
        ("Verified", f'=COUNTIF({rng("Outcome")},"Verified")', n_ver),
        ("Held for human", f'=COUNTIF({rng("Outcome")},"Held for human")', n_held),
        ("Failed", f'=COUNTIF({rng("Outcome")},"Failed")', n_fail),
        ("Withdrawn", f'=COUNTIF({rng("Outcome")},"Withdrawn")', n_canc),
        ("Moved to Approved", f'=COUNTIF({rng("Moved to Approved")},"Yes")', n_appr),
        ("On Shopify now", f'=COUNTIF({rng("On Shopify")},"Yes")', None),
        ("Verified with rules still open (approved; worth a look)",
         f'=COUNTIFS({rng("Outcome")},"Verified",{rng("Rules still open")},"<>")', None),
        (None, None, None),
        ("Why not verified", "Products", None),
    ]
    # Held rows grouped by their codes (short, exact); failed rows by their sentence.
    groups: list[tuple[str, str, str]] = []
    for t in table:
        if t[0] == "Held for human":
            g = (t[0], f'{rng("Rules still open")},"{t[5]}",{rng("Preflight problems")},"{t[6]}"',
                 " + ".join(filter(None, [t[5], t[6]])))
        elif t[0] == "Failed":
            g = (t[0], f'{rng("Why")},"{t[4]}"', t[4])
        else:
            continue
        if g not in groups:
            groups.append(g)
    for o, crit, label in groups:
        lines.append((f"{o}: {label}", f'=COUNTIFS({rng("Outcome")},"{o}",{crit})', None))
    lines += [(None, None, None), ("What the codes mean", None, None)]
    seen = sorted({c for t in table for c in (t[5].split(", ") + t[6].split(", ")) if c})
    lines += [(code, MEANING.get(code, ""), None) for code in seen]
    lines += [(None, None, None),
              ("Source", f"{db}, read-only, at {read_at.astimezone(tz):%Y-%m-%d %H:%M} {zone}", None),
              ("Note", "Title at run start is the run's own snapshot; the 'now' columns are the "
                       "product as it is today.", None)]
    section = {"Run", "Products", "Why not verified", "What the codes mean"}
    for a, b, cc in lines:
        summary.append([a, b, cc])
        row = summary.max_row
        for cell in summary[row]:
            cell.font = font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if a in section:
            for cell in summary[row]:
                cell.font, cell.fill = head, head_fill
        if isinstance(b, datetime):
            summary.cell(row, 2).number_format = "yyyy-mm-dd hh:mm"
    summary.column_dimensions["B"].width = 60

    wb.calculation.fullCalcOnLoad = True
    args.out.mkdir(parents=True, exist_ok=True)
    out = args.out / f"{tenant.lower()}_run_{local(started):%Y-%m-%d_%H%M}_{run_id[:8]}.xlsx"
    wb.save(out)
    counts = {o: sum(1 for t in table if t[0] == o) for o in ORDER if any(t[0] == o for t in table)}
    print(f"{tenant} run {run_id[:8]} ({db}): {len(table)} products {counts}; "
          f"run counters verified={n_ver} held={n_held} failed={n_fail} approved={n_appr}")
    print(f"  {out}")


if __name__ == "__main__":
    main()
