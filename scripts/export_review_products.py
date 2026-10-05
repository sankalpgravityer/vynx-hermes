#!/usr/bin/env python
"""Export one tenant's Review-tab products to Excel + a SKU list. READ ONLY.

    python scripts/export_review_products.py --tenant BOAS
    python scripts/export_review_products.py --tenant BOAS --generation COMPLETE   (the default)
    python scripts/export_review_products.py --tenant BOAS --generation ANY

The Review tab is `currentStage = 'REVIEW'` (the stage is derived and stored on
the row). The database comes from Hermes' .env DATABASE_URL unless --db is
given, and the connection is opened read-only.

Writes reports/<tenant>_review_<generation>_<date>.xlsx and ..._skus.txt. The
workbook's first sheet has `Product id` and `SKU` columns, so it can be handed
straight to scripts/run_from_sheet.py --sheet.
"""
from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
UI = "https://try.vnyx.ai"

COLUMNS = [  # (header, sql expression, width)
    ("Product id", 'p.id::text', 38),
    ("SKU", 'p.sku', 14),
    ("Title", 'p.title', 60),
    ("Brand", 'b.name', 20),
    ("Master category", 'p."masterCategory"', 16),
    ("Category", 'p.category', 18),
    ("Subcategory", 'p."subCategory"', 20),
    ("Size", 'p."internationalSize"', 10),
    ("Price", 'p.price', 10),
    ("Retail price", 'p."retailPrice"', 12),
    ("Status", 'p.status::text', 10),
    ("Review status", 'p."reviewStatus"::text', 16),
    ("Generation", 'p."generationStatus"::text', 13),
    ("Verification", 'p."verificationStatus"::text', 14),
    ("Created (UTC)", 'p."createdAt"', 18),
    ("In Review since (UTC)", 'p."currentStageAt"', 20),
]


def dsn_from_env() -> str:
    env = (ROOT / ".env").read_text(encoding="utf-8")
    m = re.search(r'^\s*DATABASE_URL\s*=\s*"?([^"\s]+)', env, re.M)
    if not m:
        raise SystemExit("no DATABASE_URL in Hermes' .env — pass --db")
    return m.group(1)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tenant", required=True, help="tenant name, e.g. BOAS")
    ap.add_argument("--generation", default="COMPLETE",
                    help="generationStatus to keep, or ANY (default COMPLETE)")
    ap.add_argument("--db", help="connection string (default: Hermes .env DATABASE_URL)")
    ap.add_argument("--out", type=Path, default=ROOT / "reports")
    args = ap.parse_args()

    gen = args.generation.upper()
    read_at = datetime.now(timezone.utc)
    with psycopg.connect(args.db or dsn_from_env(), connect_timeout=30,
                         options="-c default_transaction_read_only=on -c statement_timeout=120000") as c:
        c.read_only = True
        t = c.execute('SELECT id::text, name FROM "Tenant" WHERE name ILIKE %s', (args.tenant,)).fetchall()
        if len(t) != 1:
            raise SystemExit(f"tenant {args.tenant!r} matched {len(t)} rows: {t}")
        tenant_id, tenant = t[0]
        by_gen = dict(c.execute(
            '''SELECT "generationStatus"::text, count(*) FROM "Product"
                WHERE "tenantId" = %s::uuid AND "currentStage" = 'REVIEW' GROUP BY 1''',
            (tenant_id,)).fetchall())
        where = '' if gen == "ANY" else 'AND p."generationStatus"::text = %(gen)s'
        rows = c.execute(
            f'''SELECT {", ".join(expr for _, expr, _ in COLUMNS)}
                  FROM "Product" p LEFT JOIN "Brand" b ON b.id = p."brandId"
                 WHERE p."tenantId" = %(t)s::uuid AND p."currentStage" = 'REVIEW' {where}
                 ORDER BY p."createdAt" DESC''',
            {"t": tenant_id, "gen": gen}).fetchall()

    args.out.mkdir(parents=True, exist_ok=True)
    stem = f"{tenant.lower()}_review_{gen.lower()}_{read_at:%Y-%m-%d}"
    xlsx, txt = args.out / f"{stem}.xlsx", args.out / f"{stem}_skus.txt"

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F4E5A")
    link = Font(name="Arial", size=10, color="0563C1", underline="single")

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = [h for h, _, _ in COLUMNS] + ["Days in Review", "Edit link"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(vertical="center")
    since_col = get_column_letter(headers.index("In Review since (UTC)") + 1)
    for i, r in enumerate(rows, start=2):
        values = [(v.replace(tzinfo=None) if isinstance(v, datetime) else v) for v in r]
        ws.append(values + [f"=IF({since_col}{i}=\"\",\"\",INT(NOW()-{since_col}{i}))",
                            f"{UI}/product/{r[0]}/edit?tenantId={tenant_id}"])
        for cell in ws[i]:
            cell.font = font
        for h in ("Created (UTC)", "In Review since (UTC)"):
            ws.cell(i, headers.index(h) + 1).number_format = "yyyy-mm-dd hh:mm"
        lc = ws.cell(i, len(headers))
        lc.hyperlink, lc.font = lc.value, link
    widths = [w for _, _, w in COLUMNS] + [14, 30]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"

    s = wb.create_sheet("Summary")
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 16
    s.append(["Item", "Value"])
    for cell in s[1]:
        cell.font, cell.fill = bold, head_fill
    last = len(rows) + 1
    summary = [
        ("Tenant", tenant),
        ("Tab", "Review (currentStage = REVIEW)"),
        ("Filter", "generationStatus = " + gen if gen != "ANY" else "none"),
        ("Products in this file", f"=COUNTA(Products!B2:B{max(2, last)})"),
        ("Of which never verified (NOT_STARTED)",
         f'=COUNTIF(Products!{get_column_letter(headers.index("Verification") + 1)}2:'
         f'{get_column_letter(headers.index("Verification") + 1)}{max(2, last)},"NOT_STARTED")'),
    ]
    for k in sorted(by_gen):
        summary.append((f"Review tab, generation {k}", by_gen[k]))
    summary += [
        ("Review tab, all generations", sum(by_gen.values())),
        ("Source", f"vnyx database, read-only, at {read_at:%Y-%m-%d %H:%M} UTC"),
    ]
    for k, v in summary:
        s.append([k, v])
    for row in s.iter_rows(min_row=2):
        for cell in row:
            cell.font = font

    # openpyxl stores no cached formula results; make Excel compute them on open.
    wb.calculation.fullCalcOnLoad = True
    wb.save(xlsx)
    txt.write_text("\n".join(r[1] for r in rows if r[1]) + "\n", encoding="utf-8")
    print(f"{tenant}: {len(rows)} product(s) in Review with generation {gen} "
          f"(Review tab total {sum(by_gen.values())}: {by_gen})")
    print(f"  {xlsx}\n  {txt}")


if __name__ == "__main__":
    main()
