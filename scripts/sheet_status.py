#!/usr/bin/env python
"""Where the products of a run sheet stand NOW — read-only.

    python scripts/sheet_status.py --sheet "reports/recheck 12 oct/recheck_klekt.xlsx"
    python scripts/sheet_status.py --sheet ... --since "2026-10-12 12:00" --out status.xlsx

The companion of run_from_sheet.py: after a run, how many of the sheet's products
are now approved, verified, still held (and for what), rejected, or not checked
yet. Counted per tenant and per the section each product was in when the sheet
was made (the sheet's `Section` column; "Review" / "Approved"), because the two
mean different things:

  * a REVIEW product has succeeded when it is APPROVED now;
  * an APPROVED product was already approved — it has succeeded when it is
    VERIFIED instead of held.

`--since` counts a product as "not checked yet" when its last verdict is older
than that time (UTC), so a half-finished run reads as half-finished. Writes
nothing to the database; `--out` writes one row per product to a workbook.
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.auto_approval import db  # noqa: E402


def read_rows(path: Path) -> list[dict[str, str]]:
    """SKU (or Product id) and Section from the first sheet that has them."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h or "").strip().lower() for h in rows[0]]
        key = next((h for h in ("sku", "product id", "productid", "id") if h in head), None)
        if key is None:
            continue
        i_key, i_sec = head.index(key), (head.index("section") if "section" in head else None)
        out = []
        for r in rows[1:]:
            v = str(r[i_key] or "").strip()
            if v:
                out.append({"key": v, "by": "sku" if key == "sku" else "id",
                            "section": str(r[i_sec] or "").strip() if i_sec is not None else ""})
        return out
    sys.exit(f"{path} has no SKU or Product id column")


def status_of(p: dict[str, Any], since: datetime | None) -> str:
    approved = p["reviewStatus"] == "ACCEPTED" or p["currentStage"] == "APPROVED"
    v = p["verificationStatus"]
    if p["reviewStatus"] == "REJECTED" or p["currentStage"] == "REJECTED":
        return "Rejected"
    if v in ("QUEUED", "IN_PROGRESS"):
        return "In the queue"
    if since and (p["checked"] is None or p["checked"] < since):
        return "Not checked yet"
    if v == "HELD_FOR_HUMAN":
        return "Still held"
    if approved and v == "VERIFIED":
        return "Approved + verified"
    if approved:
        return "Approved"
    if v == "VERIFIED":
        return "Verified, not approved"
    return v.replace("_", " ").capitalize() if v else "Unknown"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--sheet", required=True, type=Path)
    ap.add_argument("--since", help='"YYYY-MM-DD HH:MM" (UTC): older verdicts count as not checked yet')
    ap.add_argument("--out", type=Path, help="also write one row per product to this .xlsx")
    args = ap.parse_args()

    since = datetime.strptime(args.since, "%Y-%m-%d %H:%M") if args.since else None
    rows = read_rows(args.sheet)
    by = rows[0]["by"] if rows else "sku"
    found = db.fetch_all(
        f"""SELECT p.id::text AS id, p.sku, t.name AS tenant, p."reviewStatus"::text AS "reviewStatus",
                   p."currentStage"::text AS "currentStage",
                   p."verificationStatus"::text AS "verificationStatus",
                   p."verificationOutcome" AS outcome, p."verificationReason" AS reason,
                   (SELECT max(x."completedAt") FROM "AutoApprovalRunProduct" x
                     WHERE x."productId" = p.id) AS checked
              FROM "Product" p JOIN "Tenant" t ON t.id = p."tenantId"
             WHERE {'p.sku' if by == 'sku' else 'p.id::text'} = ANY(%(k)s)""",
        {"k": [r["key"] for r in rows]})
    index = {(f["sku"] if by == "sku" else f["id"]): f for f in found}

    table: dict[tuple[str, str], Counter] = defaultdict(Counter)
    held_why: Counter = Counter()
    out_rows = []
    for r in rows:
        p = index.get(r["key"])
        if p is None:
            table[("?", r["section"] or "?")]["Not found"] += 1
            continue
        if p["checked"] is not None and p["checked"].tzinfo is not None:
            p["checked"] = p["checked"].replace(tzinfo=None)
        st = status_of(p, since)
        table[(p["tenant"], r["section"] or "-")][st] += 1
        if st == "Still held":
            held_why[(p["outcome"] or "?")[:40]] += 1
        out_rows.append([p["tenant"], r["section"], p["sku"], st, p["reviewStatus"], p["currentStage"],
                         p["verificationStatus"], p["outcome"], (p["reason"] or "")[:300],
                         str(p["checked"] or "")[:16]])

    cols = ["Approved + verified", "Approved", "Verified, not approved", "Still held", "Rejected",
            "In the queue", "Not checked yet", "Not found"]
    extra = sorted({k for c in table.values() for k in c} - set(cols))
    cols += extra
    print(f"\n{args.sheet.name}: {len(rows)} product(s)\n")
    print(f"{'Tenant':8} {'Section':9} " + " ".join(f"{c[:14]:>14}" for c in cols) + f" {'Total':>6}")
    for (t, s), c in sorted(table.items()):
        print(f"{t:8} {s:9} " + " ".join(f"{c.get(k, 0):>14}" for k in cols) + f" {sum(c.values()):>6}")
    total = Counter()
    for c in table.values():
        total.update(c)
    print(f"{'Total':18} " + " ".join(f"{total.get(k, 0):>14}" for k in cols) + f" {sum(total.values()):>6}")
    if held_why:
        print("\nStill held — top reasons:")
        for k, n in held_why.most_common(12):
            print(f"  {n:5}  {k}")

    if args.out:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Status"
        ws.append(["Tenant", "Section (sheet)", "SKU", "Status now", "Review status", "Stage",
                   "Verification", "Held code", "Agent's reason", "Last checked (UTC)"])
        for row in out_rows:
            ws.append(row)
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", bold=(cell.row == 1))
        wb.save(args.out)
        print(f"\nwrote {args.out}")


if __name__ == "__main__":
    main()
