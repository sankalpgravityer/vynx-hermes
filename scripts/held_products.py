#!/usr/bin/env python
"""Every product the agent holds for a person — the list, and the run sheets. READ-ONLY.

    python scripts/held_products.py --out "reports/held run/held_products.xlsx"
    python scripts/held_products.py --tenant Klekt,BOAS --out ...

Writes, beside --out:
  * the list: every HELD_FOR_HUMAN product (not deleted, not archived), tenant by
    tenant — section, stage, the hold code and reason, the kind of problem, and
    whether `run_from_sheet.py` will take it with the flags the run uses
    (--retry --include-label --include-approved) and, if not, which of its
    eligibility clauses it fails — the runner's own (run_from_sheet.
    eligibility_clauses), so the count here is the count the run takes;
  * run_<Tenant>.xlsx per tenant: the products that go into the run. One per
    tenant because `run_from_sheet.py --workers N` verifies N products of ONE
    tenant at once.

A Rejected product is listed, never in a run sheet: approve-products.ts refuses
REJECTED as a stage to move from, and the runner's stages never include it.
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import psycopg  # noqa: E402
from psycopg.rows import dict_row  # noqa: E402

from app.services.auto_approval import db, sweep  # noqa: E402
from scripts import run_from_sheet as rfs  # noqa: E402

# The run's flags: --retry (they carry a verdict), --include-label (a Review
# product whose stage flag reads LABEL), --include-approved (the Approved section).
RUN_STAGES = ["REVIEW", "LABEL", "APPROVED"]

HELD_SQL = """
SELECT p.id::text AS id, t.name AS tenant, p."tenantId"::text AS tenant_id, p.sku, p.title,
       p."reviewStatus"::text AS review, p."currentStage"::text AS stage,
       p."generationStatus"::text AS generation, p."verificationOutcome" AS code,
       p."verificationReason" AS reason, p."verificationCheckedAt" AS checked_at,
       p."shopifyProductId" AS shopify_id,
       EXISTS (SELECT 1 FROM "AutoApprovalConfig" c WHERE c."tenantId" = p."tenantId") AS has_config,
       {checks}
  FROM "Product" p JOIN "Tenant" t ON t.id = p."tenantId"
 WHERE p."verificationStatus"::text = 'HELD_FOR_HUMAN'
   AND p."isDeleted" = false AND p."isArchived" = false
   {tenants}
 ORDER BY t.name, p."reviewStatus", p.sku
"""

# What kind of problem a hold code is, for the reader — by code or prefix.
KINDS: list[tuple[tuple[str, ...], str]] = [
    (("CUTOUT_", "PREFLIGHT_CUT_OUTS", "PREFLIGHT_PHOTO_AUDIT__CUTOUT"), "Cut-out"),
    (("NO_SIZE_CHART", "NO_SIZING_GUIDE", "SIZE."), "Size / size chart"),
    (("NO_BRAND_OR_SIZE", "BRAND_MISSING", "NO_BRAND", "NO_SIZE", "NO_EU_SIZE"), "Brand / size missing"),
    (("DATA.",), "Missing field (EU size, colour …)"),
    (("WOMEN_RIG", "MEN_RIG", "UPPER_RIG", "LOWER_RIG", "MANNEQUIN"), "Mannequin / rig"),
    (("TAX.", "MASTER_CATEGORY", "CATEGORY_IMAGE_MISMATCH", "PREFLIGHT_IMAGE_GATE", "DRIFT."),
     "Category / taxonomy"),
    (("PRICE.", "GRADE."), "Price / grade"),
    (("NO_FRONT_PHOTO", "NO_BACK_PHOTO", "NO_CARE_LABEL", "IMG."), "Photos"),
    (("MISSING_RENDER", "RENDER_", "PREFLIGHT_PHOTO_AUDIT__RENDER"), "Renders"),
    (("TEXT.",), "Title / description"),
    (("VISION_UNAVAILABLE", "LEASE_EXPIRED"), "Not judged (retry)"),
]

SECTION = {"ACCEPTED": "Approved", "REJECTED": "Rejected"}


def kind_of(code: str | None) -> str:
    c = str(code or "")
    for prefixes, name in KINDS:
        if any(c.startswith(p) for p in prefixes):
            return name
    return "Other"


def load(tenants: list[str]) -> tuple[list[dict[str, Any]], list[tuple[str, str]]]:
    clauses = rfs.eligibility_clauses(retry=True, stages=RUN_STAGES)
    checks = ",\n       ".join(f'({sql}) AS "chk_{i}"' for i, (_, sql) in enumerate(clauses))
    sql = HELD_SQL.format(checks=checks, tenants="AND t.name = ANY(%(t)s)" if tenants else "")
    with psycopg.connect(db.dsn(), row_factory=dict_row) as conn:
        conn.execute("SET SESSION CHARACTERISTICS AS TRANSACTION READ ONLY")
        rows = conn.execute(sql, {"t": tenants, "gen": sweep.generation_statuses(False)}).fetchall()
    for r in rows:
        r["section"] = SECTION.get(r["review"], "Review")
        r["kind"] = kind_of(r["code"])
        fails = [name for i, (name, _) in enumerate(clauses) if not r[f"chk_{i}"]]
        if not r["has_config"]:
            fails.append("tenant has no agent config")
        r["why_not"] = ", ".join(fails)
        r["in_run"] = not fails
    return rows, clauses


def _style(ws, widths) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for cell in ws[1]:
        cell.font, cell.fill = Font(name="Arial", bold=True), PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


HEAD = ["Tenant", "SKU", "Title", "Section", "Stage", "Generation", "Hold code", "Kind of problem",
        "Hold reason", "Held since (UTC)", "Goes in the run", "Why not", "On Shopify", "Open in VNYX",
        "Product id"]
WIDTHS = [12, 13, 40, 10, 13, 11, 30, 22, 70, 16, 10, 34, 10, 40, 38]


def _row(r: dict[str, Any]) -> list[Any]:
    return [r["tenant"], r["sku"], r["title"], r["section"], r["stage"], r["generation"], r["code"],
            r["kind"], (r["reason"] or "")[:500], str(r["checked_at"] or "")[:16],
            "yes" if r["in_run"] else "no", r["why_not"], "yes" if r["shopify_id"] else "no",
            f"https://try.vnyx.ai/product/{r['id']}/edit?tenantId={r['tenant_id']}", r["id"]]


def write(out: Path, rows: list[dict[str, Any]]) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "All held"
    ws.append(HEAD)
    for r in rows:
        ws.append(_row(r))
    _style(ws, WIDTHS)
    last = ws.max_row
    tenants = sorted({r["tenant"] for r in rows})
    for t in tenants:
        wt = wb.create_sheet(t[:31])
        wt.append(HEAD)
        for r in rows:
            if r["tenant"] == t:
                wt.append(_row(r))
        _style(wt, WIDTHS)

    s = wb.create_sheet("Summary", 0)
    bold = Font(name="Arial", bold=True)
    s["A1"] = "Products held for a person — and the Hermes run"
    s["A1"].font = Font(name="Arial", bold=True, size=13)
    s["A2"] = (f"Built {datetime.now():%d %b %Y %H:%M} from production, read-only. Every product whose "
               "verification is HELD_FOR_HUMAN (not deleted, not archived). 'Goes in the run' uses "
               "run_from_sheet.py's own eligibility rules with --retry --include-label --include-approved; "
               "Rejected products are never run.")
    s["A2"].font = Font(name="Arial", italic=True, size=9)
    s["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A2:H2")
    s.row_dimensions[2].height = 52
    rng = f"'All held'!$A$2:$A${last}"
    sec, run = f"'All held'!$D$2:$D${last}", f"'All held'!$K$2:$K${last}"
    head = ["Tenant", "Held", "Review", "Approved", "Rejected", "Goes in the run", "Not in the run"]
    for j, h in enumerate(head, 1):
        s.cell(row=4, column=j, value=h).font = bold
    for i, t in enumerate(tenants, 5):
        vals = [t, f"=COUNTIF({rng},A{i})",
                f'=COUNTIFS({rng},A{i},{sec},"Review")', f'=COUNTIFS({rng},A{i},{sec},"Approved")',
                f'=COUNTIFS({rng},A{i},{sec},"Rejected")', f'=COUNTIFS({rng},A{i},{run},"yes")',
                f"=B{i}-F{i}"]
        for j, v in enumerate(vals, 1):
            s.cell(row=i, column=j, value=v).font = Font(name="Arial")
    tot = 5 + len(tenants)
    s.cell(row=tot, column=1, value="Total").font = bold
    for j, c in zip(range(2, 8), "BCDEFG"):
        s.cell(row=tot, column=j, value=f"=SUM({c}5:{c}{tot - 1})").font = bold

    # by kind of problem, per tenant
    k0 = tot + 3
    s.cell(row=k0 - 1, column=1, value="By kind of problem").font = bold
    kinds = [name for _, name in KINDS] + ["Other"]
    kcol = f"'All held'!$H$2:$H${last}"
    s.cell(row=k0, column=1, value="Kind of problem").font = bold
    for j, t in enumerate(tenants, 2):
        s.cell(row=k0, column=j, value=t).font = bold
    s.cell(row=k0, column=len(tenants) + 2, value="Total").font = bold
    for i, k in enumerate(kinds, k0 + 1):
        s.cell(row=i, column=1, value=k).font = Font(name="Arial")
        for j, t in enumerate(tenants, 2):
            s.cell(row=i, column=j, value=f'=COUNTIFS({kcol},$A{i},{rng},"{t}")').font = Font(name="Arial")
        last_t = s.cell(row=i, column=len(tenants) + 1).column_letter
        s.cell(row=i, column=len(tenants) + 2,
               value=f"=SUM(B{i}:{last_t}{i})").font = Font(name="Arial")
    for c, w in zip("ABCDEFGH", (34, 10, 10, 10, 10, 14, 14, 10)):
        s.column_dimensions[c].width = w
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)

    # the run sheets, one per tenant
    written = []
    for t in tenants:
        mine = [r for r in rows if r["tenant"] == t and r["in_run"]]
        if not mine:
            continue
        rw = Workbook()
        w = rw.active
        w.title = "Products"
        w.append(["Product id", "SKU", "Tenant", "Section", "Hold code", "Kind of problem"])
        for r in mine:
            w.append([r["id"], r["sku"], r["tenant"], r["section"], r["code"], r["kind"]])
        _style(w, [38, 13, 12, 10, 30, 22])
        p = out.parent / f"run_{t.replace(' ', '_')}.xlsx"
        rw.save(p)
        written.append(p)
        print(f"wrote {p} — {len(mine)} product(s)")
    return written


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--tenant", help="comma list of tenant names (default: every tenant)")
    ap.add_argument("--out", type=Path, required=True, help="the list; the run sheets go beside it")
    args = ap.parse_args()
    tenants = [t.strip() for t in (args.tenant or "").split(",") if t.strip()]
    rows, _ = load(tenants)
    by_t = Counter(r["tenant"] for r in rows)
    run_t = Counter(r["tenant"] for r in rows if r["in_run"])
    print(f"{len(rows)} held product(s); {sum(run_t.values())} go in the run")
    for t in sorted(by_t):
        print(f"  {t:20} held {by_t[t]:5}   in the run {run_t[t]:5}")
    why = Counter(r["why_not"] for r in rows if not r["in_run"])
    for w, n in why.most_common():
        print(f"  not in the run ({n}): {w}")
    write(args.out, rows)


if __name__ == "__main__":
    main()
