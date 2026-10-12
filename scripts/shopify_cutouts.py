#!/usr/bin/env python
"""The cut-outs approved products send to Shopify — list, clean, check, push.

Scope: APPROVED products of the named tenants that the agent has verified at least
once (verificationStatus is anything but NOT_STARTED) and that are on Shopify.
Cut-outs only — no size chart, no care label.

    # 1. the sheet, read-only — no AI, no writes
    python scripts/shopify_cutouts.py --out "reports/shopify cutouts/to_fix.xlsx"
    python scripts/shopify_cutouts.py --tenant Klekt --out ...
    python scripts/shopify_cutouts.py --skus BOA-000314,KLE-000536 --out ...

    # 2. the fix (WRITES + pushes) — the products the sheet lists
    python scripts/shopify_cutouts.py --sheet "reports/shopify cutouts/to_fix.xlsx" --apply \
        --out "reports/shopify cutouts/fix_results.xlsx"
    ... --limit 10        first N products only (try it on a few first)
    ... --workers 4       products in parallel (default 3)
    ... --no-push         clean and check, but do not push to Shopify
    ... --no-ai-check     clean and push, no AI look at the cut-outs
    ... --all             every product in scope, not only the ones with something to remove
                          (the AI look for all; only a changed gallery is pushed)
    ... --model NAME      the Gemini model for the cut-out check

WHAT IS WRONG ON SHOPIFY (app/imaging/gallery_cleanup.py), cut-outs only:
  * decision-portal cut-outs / photos. The Image Sequence's decision slots are off,
    so they are no longer in what a push sends; the push takes them off Shopify.
    Nothing is deleted — they stay in VNYX as internal photos.
  * duplicate cut-outs — the old pipeline's cut-out left beside a re-cut of the
    same photo (BOA-000314), or two live cut-outs of one photo. Soft-deleted: the
    row and the file stay, and vnyx-api's restoreImages undoes it.

THE FIX, per product, in this order:
  1. duplicates soft-deleted, then the gallery cache rebuilt through vnyx-api's
     `reorder` step (rebuild-media-cache.ts) — the same step the repair chain uses;
  2. Gemini compares each cut-out the listing will show with the photograph it was
     cut from, the top enlarged where the hanger and clips are
     (app/imaging/cutout_integrity) — fabric missing or cut away (a collar taken
     with the hanger hook, a waistband notched by the clips), a broken outline,
     background left in, an unclear cut-out. A hanger or clip stub alone is noted,
     never failed (keep the hanger rather than cut cloth);
  3. the push, through the `resync` step (resync-listings.ts --include-synced) —
     queued on the backend behind VNYX_API_URL, whose Shopify worker sends it —
     when the gallery changed;
  4. a cut-out the AI calls defective puts the product's verification to
     HELD_FOR_HUMAN, outcome CUTOUT_DEFECT, the reason naming the cut-out defect,
     and logs a `hold` event — so a person re-does the cut-out. The listing stays live.
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import os
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import psycopg  # noqa: E402
from psycopg.rows import dict_row  # noqa: E402

from app.config import policy  # noqa: E402
from app.imaging import gallery_cleanup as gc  # noqa: E402
from app.services.auto_approval import db, events  # noqa: E402

SCOPE_SQL = """
SELECT p.id::text AS id, t.name AS tenant, p."tenantId"::text AS tenant_id, p.sku, p.title,
       p."verificationStatus"::text AS vs, p.images, p.category, p."subCategory" AS subcategory,
       p."shopifyProductId" AS shopify_id,
       (SELECT max(l."lastSyncAt") FROM "MarketplaceListing" l WHERE l."productId" = p.id) AS last_sync
  FROM "Product" p JOIN "Tenant" t ON t.id = p."tenantId"
 WHERE t.name = ANY(%(tenants)s)
   AND p."isDeleted" = false AND p."isArchived" = false
   AND p."reviewStatus"::text = 'ACCEPTED'
   AND p."verificationStatus"::text <> 'NOT_STARTED'
   AND p."shopifyProductId" IS NOT NULL
   {skus}
 ORDER BY t.name, p.sku
"""

MEDIA_SQL = """
SELECT id::text, "productId"::text AS pid, url, view::text, origin::text, processing::text,
       "derivedFromId"::text AS "derivedFromId", "deletedAt", "createdAt", width, height, position
  FROM "ProductMedia"
 WHERE "productId"::text = ANY(%(ids)s) AND "mediaType"::text = 'IMAGE'
"""

_print_lock = threading.Lock()


def say(*a: Any) -> None:
    with _print_lock:
        print(*a, flush=True)


def read_sheet_skus(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h or "").strip().lower() for h in rows[0]]
        if "sku" in head:
            i = head.index("sku")
            return list(dict.fromkeys(str(r[i]).strip() for r in rows[1:] if r[i]))
    sys.exit(f"{path} has no SKU column")


def load(dsn: str, tenants: list[str], skus: list[str], limit: int):
    with psycopg.connect(dsn, row_factory=dict_row) as conn:
        conn.execute("SET SESSION CHARACTERISTICS AS TRANSACTION READ ONLY")
        products = conn.execute(
            SCOPE_SQL.format(skus="AND p.sku = ANY(%(skus)s)" if skus else ""),
            {"tenants": tenants, "skus": skus}).fetchall()
        if limit:
            products = products[:limit]
        media: dict[str, list[dict[str, Any]]] = defaultdict(list)
        ids = [p["id"] for p in products]
        for i in range(0, len(ids), 500):
            for r in conn.execute(MEDIA_SQL, {"ids": ids[i:i + 500]}).fetchall():
                media[r["pid"]].append(r)
    return products, media


# --------------------------------------------------------------------------- #
# the AI look at the cut-outs
# --------------------------------------------------------------------------- #
AI_MODEL: str | None = None   # --model; None = llm.model_fast


def ai_check(product: dict[str, Any], cutouts: list[dict[str, Any]],
             by_id: dict[str, dict[str, Any]]) -> tuple[list[str], str, list[str]]:
    """Gemini compares each cut-out the listing shows with the photograph it was cut
    from (app/imaging/cutout_integrity) — the whole picture and its top enlarged,
    where the hanger and the clips are.

    Returns (defects, note, seen). `defects` fail the product: garment fabric
    missing or cut away, a broken outline, background or an unclear cut-out the AI
    calls clear. `seen` is everything else it noticed (a hanger or clip stub) —
    reported, failing nothing. `note` says when the AI could not answer, in which
    case nothing is concluded.

    It replaced the photo audit here (12 Oct 2026): judging each cut-out alone, the
    audit passed BOA-002833's waistband notched by the clips and BOA-002039's
    collar cut away round the hanger hook.
    """
    from app.imaging import cutout_integrity as ci

    if not cutouts:
        return [], "no cut-out to look at", []
    pairs = [{"view": r["view"], "cutout": r["url"],
              "original": gc.root_of(r, by_id)["url"] if r.get("derivedFromId") else None}
             for r in cutouts]
    category = " > ".join(x for x in (product.get("category"), product.get("subcategory")) if x)
    answer, note = ci.judge(pairs, category=category or None, model=AI_MODEL)
    if answer is None:
        return [], note or "AI check could not run", []
    return ci.failing(answer), note, ci.noted(answer)


# --------------------------------------------------------------------------- #
# the fix, one product
# --------------------------------------------------------------------------- #
def fix(product: dict[str, Any], cls: dict[str, Any], *, dsn: str, push: bool, ai: bool) -> dict[str, Any]:
    from app.rules import imagery
    from scripts import repair_product as rp

    res: dict[str, Any] = {"done": [], "defects": [], "ai_note": "", "seen": [], "status": ""}
    vnyx_dir = Path(os.getenv("VNYX_API_DIR") or ".")
    dupes = cls["duplicates"]
    # The gallery cache still listing a row that is removed or replaced — what a
    # duplicate removal leaves until the rebuild, and what an interrupted earlier
    # fix leaves for good. Only rows that exist count: a legacy cache entry with no
    # row behind it is not this script's to drop.
    rows = list(cls["by_id"].values())
    live_urls = {r["url"] for r in gc.live(rows)}
    stale = [u for u in (product["images"] or [])
             if u not in live_urls and any(r["url"] == u for r in rows)]

    # 1. duplicates — the soft delete vnyx-api's removeImages does, then the cache.
    if dupes:
        n = db.execute(
            'UPDATE "ProductMedia" SET "deletedAt" = now(), "deletedById" = NULL'
            ' WHERE "productId" = %(p)s::uuid AND id::text = ANY(%(ids)s) AND "deletedAt" IS NULL',
            {"p": product["id"], "ids": [d["id"] for d in dupes]})
        res["done"].append(f"removed {n} duplicate cut-out(s)")
    if dupes or stale:
        # A gallery flagged as arranged by a person is rebuilt only when policy says
        # the flag is not trusted — the repair chain's own rule (order step).
        respect = bool(imagery.gallery_config(policy()).get("respect_manual", True))
        ok, _, _ = rp.run_step(vnyx_dir, "rebuild-media-cache.ts",
                               ["--db", dsn, "--product", product["id"],
                                *([] if respect else ["--include-manual"]), "--apply"],
                               timeout_s=120, quiet=True)
        res["done"].append("gallery rebuilt" if ok else "GALLERY REBUILD FAILED")

    # 2. the AI look at what the listing will show.
    if ai:
        res["defects"], res["ai_note"], res["seen"] = ai_check(product, cls["sent_cutouts"], cls["by_id"])
        res["done"].append(f"AI: {len(res['defects'])} defect(s)" if not res["ai_note"].startswith("AI check could not")
                           else res["ai_note"])

    # 3. the push — only when the listing's gallery actually changes: a product run
    # for the AI look alone (--all) has nothing new to send.
    if push and (dupes or stale or cls["decision"]):
        ok, out, _ = rp.run_step(vnyx_dir, "resync-listings.ts",
                                 ["--product", product["id"], "--include-synced", "--apply"],
                                 timeout_s=120, quiet=True)
        m = re.search(r"queued\s*:\s*(\d+)", out or "")
        if ok and m and int(m.group(1)) > 0:
            res["done"].append("Shopify push queued")
        else:
            why = [ln.split("skipped —", 1)[1].strip() for ln in (out or "").splitlines() if "skipped —" in ln]
            res["done"].append("push NOT queued" + (f" — {'; '.join(why)}" if why else ""))

    # 4. a defective cut-out holds the product for a person.
    if res["defects"]:
        # Each finding already reads "CUTOUT DEFECT — <view> cut-out: …".
        reason = "; ".join(res["defects"])
        db.execute(
            """UPDATE "Product"
                  SET "verificationStatus" = 'HELD_FOR_HUMAN'::"VerificationStatus",
                      "verificationOutcome" = 'CUTOUT_DEFECT',
                      "verificationReason" = %(r)s,
                      "verificationCheckedAt" = now()
                WHERE id = %(p)s::uuid""",
            {"p": product["id"], "r": reason[:500]})
        events.emit(product["tenant_id"], "hold",
                    f"{product['sku']} → held for human · CUTOUT_DEFECT (Shopify cut-out check)",
                    product_id=product["id"],
                    detail={"defects": res["defects"], "by": "shopify_cutouts.py"})
        res["status"] = "HELD_FOR_HUMAN"
    else:
        res["status"] = product["vs"]
    return res


# --------------------------------------------------------------------------- #
# workbooks
# --------------------------------------------------------------------------- #
def _style(ws, widths) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="DDE4EE")
    for cell in ws[1]:
        cell.font, cell.fill = Font(name="Arial", bold=True), head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def write_list(out: Path, products, classes) -> None:
    """The sheet to check before the fix: what each product has that Shopify should not."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wp = wb.active
    wp.title = "Products"
    wp.append(["Tenant", "SKU", "Title", "Verification", "Decision cut-outs (leave Shopify)",
               "Duplicate cut-outs (removed)", "Cut-outs Shopify keeps", "What the fix does",
               "Decision files", "Duplicate files", "Shopify product id", "Last Shopify sync (UTC)"])
    wr = wb.create_sheet("Images")
    wr.append(["Tenant", "SKU", "What it is", "What the fix does", "View", "Origin", "Processing", "Size",
               "Created", "In what a push sends now", "File", "Cut from (original photo)", "URL", "Media id"])
    listed = 0
    for p in products:
        c = classes[p["id"]]
        if not (c["decision"] or c["duplicates"]):
            continue
        listed += 1
        wp.append([p["tenant"], p["sku"], p["title"], p["vs"], len(c["decision"]), len(c["duplicates"]),
                   len(c["sent_cutouts"]),
                   "; ".join(x for x in ("remove the duplicates" if c["duplicates"] else "",
                                         "AI-check the cut-outs", "re-push to Shopify") if x),
                   ", ".join(gc.base(x["url"]) for x in c["decision"]),
                   ", ".join(gc.base(x["url"]) for x in c["duplicates"]),
                   str(p["shopify_id"]), str(p["last_sync"] or "")[:16]])
        sent = set(p["images"] or [])
        for kind, rows_ in (("decision-portal cut-out / photo", c["decision"]),
                            ("duplicate — old pipeline cut-out of a re-cut photo", c["legacy_duplicates"]),
                            ("duplicate — second cut-out of the same photo", c["same_photo_duplicates"])):
            for x in rows_:
                src = gc.root_of(x, c["by_id"])
                wr.append([p["tenant"], p["sku"], kind,
                           ("stays in VNYX (hidden by the image sequence) — the push takes it off Shopify"
                            if kind.startswith("decision") else "soft-deleted — the push takes it off Shopify"),
                           x["view"], x["origin"], x["processing"],
                           f"{x['width']}x{x['height']}" if x.get("width") else "",
                           str(x.get("createdAt") or "")[:10], "yes" if x["url"] in sent else "no",
                           gc.base(x["url"]), gc.base(src["url"]) if x.get("derivedFromId") else "",
                           x["url"], x["id"]])
    _style(wp, [9, 13, 44, 16, 13, 13, 11, 44, 50, 50, 16, 17])
    _style(wr, [9, 13, 40, 56, 7, 11, 11, 11, 11, 11, 44, 44, 60, 38])

    s = wb.create_sheet("Summary", 0)
    s["A1"] = "Cut-outs to fix on Shopify — approved products the agent has verified"
    s["A1"].font = Font(name="Arial", bold=True, size=13)
    s["A2"] = (f"Built {datetime.now():%d %b %Y %H:%M} from production, read-only. Approved products on "
               "Shopify whose verification status is not 'not started'. Cut-outs only. Each image is classified "
               "by the photo it was cut from, not its file name or size. The fix removes the duplicates, "
               "AI-checks the cut-outs Shopify keeps (a defective one holds the product for a person, "
               "reason: cut-out defect), and re-pushes; decision cut-outs stay in VNYX and leave Shopify "
               "with the push.")
    s["A2"].font = Font(name="Arial", italic=True, size=9)
    s["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A2:F2")
    s.row_dimensions[2].height = 64
    hdr = ["Tenant", "In scope", "With decision cut-outs", "With duplicates", "To fix", "Nothing to fix"]
    for j, h in enumerate(hdr, 1):
        s.cell(row=4, column=j, value=h).font = Font(name="Arial", bold=True)
    tenants = sorted({p["tenant"] for p in products})
    for i, t in enumerate(tenants, 5):
        mine = [p for p in products if p["tenant"] == t]
        dec = sum(1 for p in mine if classes[p["id"]]["decision"])
        dup = sum(1 for p in mine if classes[p["id"]]["duplicates"])
        fixn = sum(1 for p in mine if classes[p["id"]]["decision"] or classes[p["id"]]["duplicates"])
        for j, v in enumerate([t, len(mine), dec, dup, fixn, len(mine) - fixn], 1):
            s.cell(row=i, column=j, value=v).font = Font(name="Arial")
    tot = 5 + len(tenants)
    s.cell(row=tot, column=1, value="Total").font = Font(name="Arial", bold=True)
    for j, col in zip(range(2, 7), "BCDEF"):
        s.cell(row=tot, column=j, value=f"=SUM({col}5:{col}{tot - 1})").font = Font(name="Arial", bold=True)
    for col, w in zip("ABCDEF", (12, 10, 14, 14, 10, 12)):
        s.column_dimensions[col].width = w
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out} — {listed} product(s) to fix")


def write_results(out: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Tenant", "SKU", "Verification before", "Verification now", "Cut-out defects (AI) — fail",
               "Also seen (AI) — not failing", "Duplicates removed", "Decision cut-outs", "What was done",
               "AI note"])
    for r in rows:
        ws.append([r["tenant"], r["sku"], r["vs"], r["status"], "; ".join(r["defects"]),
                   "; ".join(r.get("seen") or []), r["dupes"], r["decision"], "; ".join(r["done"]),
                   r["ai_note"]])
    _style(ws, [9, 13, 16, 16, 70, 50, 10, 10, 60, 30])
    for row in ws.iter_rows(min_row=2):
        if row[3].value == "HELD_FOR_HUMAN" and row[4].value:
            row[3].fill = PatternFill("solid", fgColor="F8CBAD")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("wrote", out)


# --------------------------------------------------------------------------- #
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--tenant", default="Klekt,BOAS,Midtex", help="comma list of tenant names")
    ap.add_argument("--skus", help="comma list — only these")
    ap.add_argument("--sheet", type=Path, help="only the SKUs of this workbook (the list you checked)")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--out", type=Path, help="the list (read-only run) or the results (--apply)")
    ap.add_argument("--apply", action="store_true", help="WRITE: fix, AI-check, push")
    ap.add_argument("--all", action="store_true", help="with --apply: every product in scope")
    ap.add_argument("--no-push", action="store_true")
    ap.add_argument("--no-ai-check", action="store_true")
    ap.add_argument("--workers", type=int, default=3)
    ap.add_argument("--model", help="the Gemini model for the cut-out check (default llm.model_fast)")
    ap.add_argument("--yes", action="store_true", help="skip the confirmation")
    args = ap.parse_args()
    global AI_MODEL
    AI_MODEL = args.model

    dsn = db.dsn()
    skus = [s.strip() for s in (args.skus or "").split(",") if s.strip()]
    if args.sheet:
        skus = [*skus, *read_sheet_skus(args.sheet)]
    tenants = [t.strip() for t in args.tenant.split(",") if t.strip()]
    products, media = load(dsn, tenants, skus, 0)
    classes = {p["id"]: gc.classify(media[p["id"]], p["images"] or []) for p in products}
    by_t = Counter(p["tenant"] for p in products)
    print(f"{len(products)} approved, verified-at-least-once product(s) on Shopify "
          f"({', '.join(f'{t} {n}' for t, n in sorted(by_t.items()))})")
    def stale(p: dict[str, Any]) -> bool:
        rows = list(classes[p["id"]]["by_id"].values())
        live_urls = {r["url"] for r in gc.live(rows)}
        return any(u not in live_urls and any(r["url"] == u for r in rows) for u in (p["images"] or []))

    todo = [p for p in products if args.all or classes[p["id"]]["decision"]
            or classes[p["id"]]["duplicates"] or stale(p)]
    if args.limit:
        todo = todo[: args.limit]
    print(f"with decision cut-outs: {sum(1 for p in products if classes[p['id']]['decision'])} · "
          f"with duplicates: {sum(1 for p in products if classes[p['id']]['duplicates'])} · "
          f"to fix: {len(todo)}")

    if not args.apply:
        if args.out:
            write_list(args.out, products, classes)
        print("read-only — nothing written. Add --apply to fix.")
        return

    # ---- the fix -------------------------------------------------------------
    from app.services.auto_approval import preflight
    from scripts import repair_product as rp

    # THE SAME DATABASE ON BOTH SIDES, before anything is written — run_from_sheet's
    # own preflight. The rebuild and the push run on the backend behind
    # VNYX_API_URL against ITS database; a product id from another database there
    # is a silent no-op at best.
    pf = preflight.check()
    print(f"transport     : {pf.get('transport')}")
    print(f"our database  : {pf.get('database')}")
    print(f"their database: {pf.get('remoteDatabase')}")
    if not pf["ok"]:
        for problem in pf["problems"]:
            print(f"  PROBLEM: {problem}")
        sys.exit("refused — fix the problems above first")

    push, ai = not args.no_push, not args.no_ai_check
    if push and rp.vnyx_api_url() and not rp.remote_supports("includeSynced"):
        sys.exit("the vnyx-api behind VNYX_API_URL cannot re-push a live listing (no includeSynced) — deploy it")
    print(f"\nFIX: {len(todo)} product(s) · duplicates removed · "
          f"{'AI cut-out check' if ai else 'no AI check'} · {'pushed to Shopify' if push else 'NOT pushed'}"
          f" · steps via {rp.vnyx_api_url() or 'the local vnyx-api checkout'}")
    if not args.yes and input("  Type APPLY to continue: ").strip() != "APPLY":
        sys.exit("not applied")

    rows: list[dict[str, Any]] = []
    t0 = time.time()

    def one(p: dict[str, Any]) -> dict[str, Any]:
        c = classes[p["id"]]
        try:
            res = fix(p, c, dsn=dsn, push=push, ai=ai)
        except Exception as exc:  # noqa: BLE001 — one product's failure is a row, not a stop
            res = {"done": [f"FAILED: {type(exc).__name__}: {exc}"[:200]], "defects": [],
                   "ai_note": "", "seen": [], "status": p["vs"]}
        return {"tenant": p["tenant"], "sku": p["sku"], "vs": p["vs"], "dupes": len(c["duplicates"]),
                "decision": len(c["decision"]), **res}

    with cf.ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        for n, r in enumerate(pool.map(one, todo), 1):
            rows.append(r)
            say(f"  [{n}/{len(todo)}] {r['sku']}: {'; '.join(r['done'])}"
                + (f"  → HELD FOR HUMAN: {'; '.join(r['defects'])[:160]}" if r["defects"] else ""))
    held = sum(1 for r in rows if r["defects"])
    print(f"\ndone in {time.time() - t0:.0f}s — {len(rows)} product(s); {held} held for human for a cut-out defect")
    if args.out:
        write_results(args.out, rows)


if __name__ == "__main__":
    main()
