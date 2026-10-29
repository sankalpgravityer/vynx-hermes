#!/usr/bin/env python
"""Put a jersey's printed PLAYER into its title and description — one product or a sheet.

    # one product, look only (nothing written)
    python scripts/apply_players.py --sku CKIL-001459

    # one product for real — no Shopify push (testing)
    python scripts/apply_players.py --sku CKIL-001459 --apply

    # a sheet for real, pushed to Shopify
    python scripts/apply_players.py --sheet ~/players-and-shoe-models.xlsx \\
        --tenant "Kilo Kilo Vintage" --apply --sync

For each product (28 Oct 2026, Kilo Kilo Vintage):
  1. the player name / number — from the sheet's "Player name" / "Player number"
     columns when it has them, else read off the back and front photographs
     (app/llm/player_print, cached — the scan's answers are reused for free);
  2. written to `properties.player_name` / `player_number` (only when different),
     and a player name the extractor filed under `model` is cleared;
  3. the title rebuilt from the tenant's title format and the description
     rewritten — vnyx-api's regenerate-copy.ts, the chain's `copy` step;
  4. with --sync, re-pushed to Shopify — resync-listings.ts, the chain's `sync`
     step (an existing listing only; never creates one).

Only for a tenant that OWNS the player_name / player_number properties, and only
when BOTH a name and a number were read — a name alone or a number alone is left
untouched. Dry run by default. Nothing is approved or moved between stages.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import product_audit  # noqa: E402
from app.config import settings  # noqa: E402
from app.rules import title_template  # noqa: E402
from scripts import repair_product as rp  # noqa: E402

KEYS = {"player_name": "playerName", "player_number": "playerNumber"}


def alnum(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def read_sheet(path: Path) -> list[dict[str, Any]]:
    """Rows of the first sheet with a SKU or Product id column, keeping the player columns."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h or "").strip().lower() for h in rows[0]]
        if "sku" not in head and "product id" not in head:
            continue
        col = {h: i for i, h in enumerate(head)}
        out = []
        for r in rows[1:]:
            get = lambda name: (str(r[col[name]]).strip() if name in col and r[col[name]] not in (None, "") else None)
            if not (get("sku") or get("product id")):
                continue
            out.append({"sku": get("sku"), "id": get("product id"), "tenant": get("tenant"),
                        "player_name": get("player name"), "player_number": get("player number"),
                        "what": get("what to do") or ""})
        return out
    sys.exit(f"{path}: no sheet with a SKU or Product id column")


def photos(loaded: dict[str, Any]) -> list[str]:
    out = []
    for view in ("BACK", "FRONT"):
        rows = [m for m in loaded["media"] if m.get("view") == view and m.get("url") and not m.get("deletedAt")]
        rows.sort(key=lambda m: (0 if m.get("processing") == "RAW" else 1, 0 if m.get("isCurrent", True) else 1))
        if rows:
            out.append(rows[0]["url"])
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--sku", help="one product (SKU or product id)")
    src.add_argument("--sheet", help="xlsx with a SKU / Product id column (the 'Players found' sheet works as is)")
    ap.add_argument("--tenant", help="only this tenant's rows (e.g. 'Kilo Kilo Vintage')")
    ap.add_argument("--apply", action="store_true", help="write (default: dry run)")
    ap.add_argument("--sync", action="store_true", help="with --apply: re-push each product to Shopify")
    # Both a name and a number, always: a number alone is often a fashion print or a
    # team shirt, and a name alone can be a slogan ("Movement").
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--db", "--dsn", dest="db")
    ap.add_argument("--results", help="write a JSON of what happened to each product")
    args = ap.parse_args()
    if args.sync and not args.apply:
        sys.exit("--sync needs --apply")

    dsn = args.db or settings().database_url
    vnyx_api = Path(rp.DEFAULT_VNYX_API)
    print(f"mode      : {'APPLY' if args.apply else 'DRY RUN'}{' + Shopify sync' if args.sync else ''}")
    print(f"vnyx-api  : {rp.vnyx_api_url() or vnyx_api}")

    if args.sku:
        entries = [{"sku": args.sku, "id": None, "player_name": None, "player_number": None, "what": ""}]
    else:
        entries = read_sheet(Path(args.sheet).expanduser())
    if args.tenant:
        entries = [e for e in entries if not e.get("tenant") or e["tenant"] == args.tenant]
    if not args.sku:
        partial = [e for e in entries if not (e.get("player_name") and e.get("player_number"))]
        entries = [e for e in entries if e.get("player_name") and e.get("player_number")]
        if partial:
            print(f"skipping  : {len(partial)} row(s) without BOTH a player name and a number")
    if args.limit:
        entries = entries[: args.limit]
    print(f"products  : {len(entries)}\n")

    import psycopg
    from app.llm import player_print

    results: list[dict[str, Any]] = []
    ctx: dict[str, Any] = {}
    counts: dict[str, int] = {}
    for n, e in enumerate(entries, 1):
        key = e["id"] or e["sku"]
        tag = f"[{n}/{len(entries)}] {e['sku'] or e['id']}"
        res: dict[str, Any] = {"key": key}
        try:
            with psycopg.connect(dsn) as c:
                c.read_only = True
                col = "id::text" if re.fullmatch(r"[0-9a-f-]{36}", key) else "sku"
                row = c.execute(f'select id::text, "tenantId"::text from "Product" where {col} = %s', (key,)).fetchone()
                if not row:
                    raise LookupError("product not found")
                pid, tid = row
                if tid not in ctx:
                    ctx[tid] = product_audit.load_tenant_context(c.cursor(), tid)
            loaded = product_audit.load(dsn, pid)
            rec = loaded["record"]
            res.update(id=pid, sku=rec.get("sku"), tenant=rec.get("tenantName"))
            if args.tenant and rec.get("tenantName") != args.tenant:
                raise LookupError(f"tenant is {rec.get('tenantName')}, not {args.tenant}")
            owned = [k for k in KEYS if k in (rec.get("tenantPropertyKeys") or [])
                     and (rec.get("defaultPropertyList") is None or k in rec["defaultPropertyList"])]
            if not owned:
                raise LookupError("the tenant has no player_name / player_number property")

            # 1. the player
            name, number = e.get("player_name"), e.get("player_number")
            source = "sheet"
            if not (name or number):
                got = player_print.read(photos(loaded))
                if got.get("error"):
                    raise RuntimeError(got["error"])
                name, number, source = got.get("player_name"), got.get("player_number"), f'photo ({got.get("provider")})'
            if name:
                name = player_print._clean_name(name)
            if number:
                number = re.sub(r"\D", "", str(number)) or None
            if not (name and number):
                counts["not both name and number"] = counts.get("not both name and number", 0) + 1
                res.update(outcome="not_both", player_name=name, player_number=number)
                print(f"{tag}  left alone — needs both a name and a number "
                      f"(read: name={name or '-'}, number={number or '-'})")
                results.append(res)
                continue

            # 2. the plan
            want = {"player_name": name, "player_number": number}
            plan = [{"kind": "set_property", "field": k, "value": v, "reason": "PLAYER_PRINT"}
                    for k, v in want.items() if v and k in owned and str(rec.get(KEYS[k]) or "") != v]
            model = str(rec.get("model") or "").strip()
            if model and name and alnum(model) in {alnum(name), alnum(f"{name}{number or ''}")}:
                plan.append({"kind": "set_property", "field": "model", "value": "", "reason": "PLAYER_PRINT"})

            # the title the format will build, for the printout
            cs = ctx[tid]["copy_settings"]
            props = {**(rec.get("properties") or {}), **{p["field"]: p["value"] for p in plan}}
            expected = title_template.resolve(
                default_template=cs.get("title_template"), overrides=cs.get("title_overrides") or [],
                sections=None, props=props, brand_relation=rec.get("brandRelation"),
                category=rec.get("category"), sub_category=rec.get("subCategory"),
                master_category=rec.get("masterCategory"), mannequin_type=rec.get("mannequinType"),
                sizing_guide=rec.get("sizingGuide"), international_size=rec.get("internationalSize"),
                sku=rec.get("sku"))["expectedTitle"]
            res.update(player_name=name, player_number=number, source=source,
                       title_before=rec.get("title"), title_after=expected)
            print(f"{tag}  {name or '-'} {number or ''}  ({source})")
            print(f"{' ' * len(tag)}  title: {rec.get('title')}")
            print(f"{' ' * len(tag)}      -> {expected}")

            if not args.apply:
                counts["would update"] = counts.get("would update", 0) + 1
                res["outcome"] = "would_update"
                results.append(res)
                continue

            if plan:
                out = product_audit.apply_plan(dsn, pid, plan, rec.get("updatedAt"))
                if out["conflict"]:
                    raise RuntimeError(out.get("message", "changed mid-update"))
                if out["skipped"]:
                    raise RuntimeError("; ".join(s.get("why", "?") for s in out["skipped"]))

            # 3. title + description, as the chain's copy step
            ok, _, payload = rp.run_step(vnyx_api, "regenerate-copy.ts",
                                         ["--db", dsn, "--product", pid, "--apply"],
                                         timeout_s=240, quiet=True, results_name="copy.json")
            if not payload or not payload.get("ok"):
                raise RuntimeError(f"regenerate-copy failed: {(payload or {}).get('error') or (payload or {}).get('failed') or 'no result'}")
            res["title_after"] = (payload.get("title") or {}).get("after") or expected
            print(f"{' ' * len(tag)}  saved title + description: {res['title_after']}")

            # 4. Shopify, as the chain's sync step
            if args.sync:
                if not rp.shopify_product_id(dsn, pid):
                    print(f"{' ' * len(tag)}  not on Shopify — no push (a re-sync would create a listing)")
                    res["sync"] = "not_on_shopify"
                else:
                    ok, out_text, _ = rp.run_step(vnyx_api, "resync-listings.ts",
                                                  ["--product", pid, "--include-synced", "--apply"],
                                                  timeout_s=120, quiet=True)
                    m = re.search(r"queued\s*:\s*(\d+)", out_text or "")
                    res["sync"] = "queued" if ok and m and int(m.group(1)) > 0 else "NOT queued"
                    print(f"{' ' * len(tag)}  Shopify: {res['sync']}")
            counts["updated"] = counts.get("updated", 0) + 1
            res["outcome"] = "updated"
        except Exception as exc:  # noqa: BLE001 — one product's failure is reported, the rest go on
            counts["failed"] = counts.get("failed", 0) + 1
            res.update(outcome="failed", error=f"{type(exc).__name__}: {exc}"[:300])
            print(f"{tag}  FAILED — {res['error']}")
        results.append(res)
        time.sleep(0.2)

    print("\nsummary   :", " · ".join(f"{k} {v}" for k, v in counts.items()) or "nothing to do")
    if args.results:
        Path(args.results).expanduser().write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
        print("results   :", args.results)
    return 0 if not counts.get("failed") else 1


if __name__ == "__main__":
    sys.exit(main())
