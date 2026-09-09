"""Annotate a product sheet with its current live status. Read-only.

Adds five columns, each answering one question a reviewer actually asks:

  Current stage          REVIEW / APPROVED / LABEL / ...
  In approved stage?     Yes / No -- the same fact as a plain yes-or-no
  Size chart present?    whether live SIZE_CHART media exists on the product
  Sizing guide correct?  whether the guide's gender matches masterCategory
  Sub-category present?  present and selectable / stored but not offered /
                         genuinely empty

WHY "PRESENT" IS THREE STATES FOR THE SUB-CATEGORY. Sub-categories hang off a
`master > category` branch. A value that exists only on the other gender's
branch is stored on the product but never offered by the dropdown, so it reads
as empty to an operator while being non-null in the database. Collapsing that
into "present" would answer the letter of the question and mislead on the
substance.

Gender is judged conservatively: only Men and Women. Kids, Unisex and unset
master categories are never called wrong, and a guide with no gender in its
name ("Defaults", "Kids", "Shoes") carries no opinion.

Usage:
  python scripts/annotate_product_status.py --sheet <xlsx>
  python scripts/annotate_product_status.py --sheet <xlsx> --out <xlsx>
"""

import argparse
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")
try:
    import psycopg
except ImportError:
    sys.exit("psycopg is required:  pip install psycopg[binary]")

DSN = "postgresql://postgres:Copenhagen%40event1@34.7.102.18:5432/vnyx-prod"

NEW_COLUMNS = [
    "Current stage",
    "In approved stage?",
    "Size chart present?",
    "Sizing guide correct?",
    "Sub-category present?",
]
HDR_FILL = PatternFill("solid", fgColor="1F2937")
GOOD = PatternFill("solid", fgColor="E8F5E9")
BADF = PatternFill("solid", fgColor="FDECEA")
WARNF = PatternFill("solid", fgColor="FFF8E1")
BLUE = PatternFill("solid", fgColor="D9E2F3")

PLACEHOLDER = {"", "unknown", "n/a", "na", "none", "null", "-", "tbd"}


def blank(v) -> bool:
    return v is None or str(v).strip().lower() in PLACEHOLDER


def gender_of(text):
    """'Men' | 'Women' | None. Women first: the word contains 'men'."""
    s = str(text or "").lower().replace(" ", "")
    if not s:
        return None
    if "women" in s or "woman" in s:
        return "Women"
    if "men" in s or "man" in s:
        return "Men"
    return None


SQL = """
SELECT p.id::text,
       p."currentStage"::text,
       p."masterCategory", p.category, p."subCategory",
       COALESCE(ps.name, p."sizingGuide")                             AS guide,
       (SELECT count(*) FROM "ProductMedia" m
         WHERE m."productId" = p.id AND m."isCurrent"
           AND m."deletedAt" IS NULL AND m.view::text = 'SIZE_CHART') AS charts,
       COALESCE(array_length(ps."sizeChartImages", 1), 0)             AS guide_imgs,
       p."tenantId"::text
FROM "Product" p
LEFT JOIN "ProductSize" ps ON ps.id::text = p."sizingGuideId"
WHERE p.id = ANY(%s::uuid[])
"""

TAXONOMY_SQL = """
WITH RECURSIVE tree AS (
  SELECT id, "tenantId", name, "parentId", name::text AS path, 1 AS depth
  FROM "Category"
  WHERE "parentId" IS NULL AND "isActive" AND "deletedAt" IS NULL
  UNION ALL
  SELECT ch.id, ch."tenantId", ch.name, ch."parentId",
         t.path || '>' || ch.name, t.depth + 1
  FROM "Category" ch JOIN tree t ON ch."parentId" = t.id
  WHERE ch."isActive" AND ch."deletedAt" IS NULL
)
SELECT "tenantId"::text, path FROM tree WHERE depth = 3
"""


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--out")
    ap.add_argument("--db", default=DSN)
    args = ap.parse_args()

    src = Path(args.sheet).expanduser()
    out = (Path(args.out).expanduser() if args.out
           else src.with_name(
               f"{src.stem}-status-{datetime.now():%Y-%m-%d-%H%M}.xlsx"))

    wb = openpyxl.load_workbook(src)
    ids: list[str] = []
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows)
                  if r and any(str(c).strip() == "Product ID"
                               for c in r if c is not None)), None)
        if h is None:
            continue
        col = list(rows[h]).index("Product ID")
        ids += [str(r[col]).strip() for r in rows[h + 1:] if r and r[col]]
    ids = list(dict.fromkeys(ids))
    print(f"Products : {len(ids):,}")

    with psycopg.connect(args.db, connect_timeout=40) as cn, cn.cursor() as cur:
        cur.execute(SQL, (ids,))
        db = {r[0]: r for r in cur.fetchall()}
        cur.execute(TAXONOMY_SQL)
        paths: dict[str, set] = {}
        for tid, path in cur.fetchall():
            paths.setdefault(tid, set()).add(path.lower())
    print(f"Found    : {len(db):,}\n")

    def verdicts(pid: str):
        rec = db.get(pid)
        if rec is None:
            return ["not found", "No", "unknown", "unknown", "unknown"]
        (_id, stage, master, cat, sub, guide, charts, guide_imgs, tid) = rec

        approved = "Yes" if stage == "APPROVED" else "No"

        if charts:
            chart = f"Yes ({charts})"
        elif not guide:
            chart = "MISSING — no sizing guide selected"
        elif not guide_imgs:
            chart = "MISSING — the guide itself has no chart image"
        else:
            chart = "MISSING — guide has an image, none copied to the product"

        m, g = gender_of(master), gender_of(guide)
        if not guide:
            guide_ok = "no guide selected"
        elif not g:
            guide_ok = f'not judged — "{guide}" carries no gender'
        elif not m:
            guide_ok = "not judged — master category is not Men/Women"
        elif m == g:
            guide_ok = "Yes"
        else:
            guide_ok = f'NO — "{guide}" is {g}, product is {m}'

        if blank(sub):
            sub_ok = "NO — empty"
        elif master and cat and \
                f"{master}>{cat}>{sub}".lower() not in paths.get(tid, set()):
            sub_ok = f'NOT SELECTABLE — "{sub}" is not offered under {master} > {cat}'
        else:
            sub_ok = "Yes"

        return [stage, approved, chart, guide_ok, sub_ok]

    tallies = [Counter() for _ in NEW_COLUMNS]
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows)
                  if r and any(str(c).strip() == "Product ID"
                               for c in r if c is not None)), None)
        if h is None:
            continue
        id_col = list(rows[h]).index("Product ID") + 1
        start = max(i for i, c in enumerate(rows[h], start=1)
                    if c is not None) + 1

        for off, name in enumerate(NEW_COLUMNS):
            c = sheet.cell(row=h + 1, column=start + off, value=name)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(start + off)].width = (
                16 if off < 2 else 52)

        for r in range(h + 2, sheet.max_row + 1):
            raw = sheet.cell(row=r, column=id_col).value
            if not raw:
                continue
            values = verdicts(str(raw).strip())
            for off, value in enumerate(values):
                cell = sheet.cell(row=r, column=start + off, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                head = str(value).split(" —")[0].split(" (")[0]
                tallies[off][head] += 1
                if head in ("Yes", "APPROVED"):
                    cell.fill = GOOD
                elif head.startswith(("NO", "MISSING", "NOT SELECTABLE")):
                    cell.fill = BADF
                elif head in ("No",) or head.startswith("not judged"):
                    cell.fill = WARNF

    sm = wb.create_sheet("Status summary", 1)
    sm.append(["Live product status"])
    sm.cell(row=1, column=1).font = Font(bold=True, size=14)
    sm.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    sm.append(["Products", len(ids)])
    for off, name in enumerate(NEW_COLUMNS):
        sm.append([])
        sm.append([name])
        c = sm.cell(row=sm.max_row, column=1)
        c.font = Font(bold=True, size=12)
        c.fill = BLUE
        print(f"--- {name} ---")
        for k, n in tallies[off].most_common():
            sm.append([k, n])
            print(f"  {n:6,}  {k}")
        print()
    for n, w in ((1, 62), (2, 12)):
        sm.column_dimensions[get_column_letter(n)].width = w
    for row in sm.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)
    print(f"written to {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
