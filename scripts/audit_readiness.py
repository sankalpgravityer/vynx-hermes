#!/usr/bin/env python
"""Readiness audit — the approved catalog against the readiness checks, read-only.

docs/READINESS-PLAN.md §5 WP-R4 (phase 0, delivered with phase 5). Every product
in scope is loaded as the chain loads it and judged with the SAME pure functions
the chain uses: the rules (the taxonomy cascade, the gallery order, the cut-out
shape and backdrop from stored evidence, the title against the record) and, when
asked, the two picture checks — `--measure-images` downloads and measures the
cut-outs (app/imaging/cutouts), `--judge-images` puts the lead render through the
image gate (app/imaging/quality_gate, one vision call per product, cached).
Nothing is written anywhere.

Two outputs: a workbook (Summary / Products / Issues), and with `--flagged` a JSON
of product ids per FIX — the backlog's input. Decision 7: only what the audit
flags is re-matted, re-rendered, re-ordered or re-written, never the whole
catalog; the commands for each list are printed at the end.

    python scripts/audit_readiness.py --db "..." --stage APPROVED --tenant BOAS --limit 200
    python scripts/audit_readiness.py --db "..." --from 2026-09-07 --to 2026-09-08 --measure-images
    python scripts/audit_readiness.py --db "..." --product <uuid,uuid> --judge-images --measure-images
    ... --out reports/readiness.xlsx --flagged reports/readiness-flagged.json --workers 4
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import product_audit  # noqa: E402
from app.config import policy  # noqa: E402

BATCH = 200
RANK = {"high": 0, "medium": 1, "low": 2}
AI_VIEWS = ("AI_FRONT", "AI_BACK", "AI_FRONT_34", "AI_BACK_34", "AI_CLOSEUP")

# The chain's rule ids, as the audit names them and what fixes each. `fix` is
# the backlog list a product lands on: rematte / regen / reorder / copy are the
# four machine repairs; `data` is the reconcile step's (a person or the chain's
# cascade); `photo` needs a camera; `none` is informational.
RULE_MAP: dict[str, tuple[str, str, str]] = {
    "TAX.001": ("MASTER_INVALID", "high", "data"),
    "TAX.004": ("GENDER_NE_MASTER", "high", "data"),
    "GENDER.001": ("GENDER_NE_MASTER", "high", "data"),
    "TAX.002": ("CATEGORY_NOT_UNDER_MASTER", "high", "data"),
    "TAX.003": ("SUBCATEGORY_NOT_UNDER_CATEGORY", "high", "data"),
    "SIZE.010": ("GUIDE_MISSING", "high", "data"),
    "SIZE.011": ("GUIDE_GENDER_MISMATCH", "high", "data"),
    "SIZE.014": ("GUIDE_SIDE_MISMATCH", "high", "data"),
    "SIZE.013": ("GUIDE_LADDER_MISSING_SIZE", "high", "data"),
    "TAX.005": ("MANNEQUIN_NE_MASTER", "low", "data"),
    "TAX.007": ("MANNEQUIN_NE_MASTER", "low", "data"),
    "IMG.025": ("ORDER_WRONG", "medium", "reorder"),
    "IMG.026": ("CUTOUT_ZOOMED", "medium", "rematte"),
    "IMG.027": ("CUTOUT_NOT_BACKDROP", "medium", "rematte"),
    "IMG.001": ("RENDER_MISSING", "high", "regen"),
    "IMG.002": ("RENDER_MISSING", "medium", "regen"),
    "IMG.010": ("CUTOUT_MISSING", "high", "rematte"),
    "IMG.030": ("LABEL_MISSING", "high", "photo"),
    "IMG.003": ("NO_GARMENT_PHOTO", "high", "photo"),
    "TEXT.002": ("TITLE_STALE", "low", "copy"),
    "TEXT.003": ("TITLE_STALE", "medium", "copy"),
    "TEXT.006": ("TITLE_STALE", "high", "copy"),
    "TEXT.007": ("TITLE_STALE", "medium", "copy"),
    "TEXT.008": ("TITLE_STALE", "medium", "copy"),
}

MEANING: dict[str, str] = {
    "MASTER_INVALID": "the master category is not a root of the tenant's tree",
    "GENDER_NE_MASTER": "the gender property disagrees with a Men / Women master",
    "CATEGORY_NOT_UNDER_MASTER": "the category is not a child of the master category",
    "SUBCATEGORY_NOT_UNDER_CATEGORY": "the subcategory is not a child of the category",
    "GUIDE_MISSING": "no sizing guide can be derived from gender + category + size",
    "GUIDE_GENDER_MISMATCH": "the sizing guide's gender disagrees with the master",
    "GUIDE_SIDE_MISMATCH": "the sizing guide measures the other half of the body",
    "GUIDE_LADDER_MISSING_SIZE": "the sizing guide's ladder does not carry the product's size",
    "MANNEQUIN_NE_MASTER": "the mannequin rig disagrees with the master (derived, informational)",
    "ORDER_WRONG": "the gallery cache is out of the catalog order (renders → uploads → booth → portal → label → chart)",
    "CUTOUT_ZOOMED": "a cut-out is not the photograph's frame — cropped to the garment, or a different aspect",
    "CUTOUT_NOT_BACKDROP": "a cut-out is not on the tenant's backdrop",
    "CUTOUT_MISSING": "a garment view has no background-removed cut-out",
    "CUTOUT_UNMEASURED": "the cut-outs could not be downloaded to measure",
    "RENDER_MISSING": "an on-model render is missing",
    "RENDER_WRONG_GENDER": "judged now: the render's model is the other gender",
    "RENDER_BODY_MISMATCH": "judged now: the model's build is two bands from the garment's size",
    "RENDER_DEFECT": "judged now: no model, a broken face or body on the lead render",
    "RENDER_CATEGORY_MISMATCH": "judged now: the render shows a different kind of garment than the category",
    "RENDER_BUILD_ONE_BAND": "judged now: the model's build is one band from the size (soft)",
    "RENDER_UNJUDGED": "the image gate could not run on this product",
    "BODY_TYPE_UNRECORDED": "the stored render build (imageSettings.bodyType) is not what the size implies — the record, not the picture",
    "LABEL_MISSING": "no care-label photograph",
    "NO_GARMENT_PHOTO": "no garment photograph to render from",
    "SIZE_CHART_MISSING": "the size chart is not on the product, or the guide has none",
    "TITLE_STALE": "the title disagrees with, or omits, the record (TEXT rules)",
    "RULE": "another Hermes rule blocks this product today",
}

FIX_COMMANDS: dict[str, str] = {
    "rematte": "repair_product.py --db <url> --products <ids> --apply --no-render   # the matte step re-cuts with --replace",
    "regen": "repair_product.py --db <url> --products <ids> --apply                 # gate -> regen -> gate",
    "reorder": "npx tsx scripts/rebuild-media-cache.ts --db <url> --product <ids> --apply   # in vnyx-api",
    "copy": "npx tsx scripts/regenerate-copy.ts --db <url> --product <id> --apply          # in vnyx-api, one id per call",
}

# Letter sizes onto the six-build scale, for BODY_TYPE_UNRECORDED.
_LETTER: dict[str, str] = {
    "xxxs": "xs", "xxs": "xs", "xs": "xs", "s": "s", "m": "m", "l": "l",
    "xl": "xl", "xxl": "xxl", "2xl": "xxl", "xxxl": "xxl", "3xl": "xxl", "4xl": "xxl",
}


def assess_readiness(record: dict[str, Any], media: list[dict[str, Any]],
                     chart: dict[str, Any] | None, findings: list[dict[str, Any]], *,
                     measured: dict[str, Any] | None = None,
                     judged: dict[str, Any] | None = None) -> list[dict[str, str]]:
    """The decision table, cold: rules + evidence -> issues with a fix each.

    `findings` are run_gate's blocking + advisory dicts; `measured` is a
    CutoutVerdict.as_dict(); `judged` a GateVerdict.as_dict(). Plain dicts in,
    plain dicts out, so the whole table is testable without a database.
    """
    issues: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    def add(code: str, severity: str, detail: str, fix: str) -> None:
        key = (code, detail[:120])
        if key in seen:
            return
        seen.add(key)
        issues.append({"code": code, "severity": severity, "detail": detail, "fix": fix})

    for f in findings:
        rid = str(f.get("rule_id") or "")
        message = str(f.get("message") or "")
        if rid in RULE_MAP:
            code, sev, fix = RULE_MAP[rid]
            add(code, sev, f"{rid}: {message}", fix)
        elif str(f.get("severity") or "").lower() in ("high", "critical"):
            add(f"RULE:{rid}", "high", message, "data")

    state = str((chart or {}).get("state") or "")
    if state and state != "ok":
        add("SIZE_CHART_MISSING", "medium", state, "data")

    # The record's own render build against the size it now carries. Renders
    # present, a letter size, a stored value: three things that must all be so.
    stored = str(((record.get("imageSettings") or {}).get("bodyType")) or "").strip().lower()
    size = str(record.get("size") or record.get("internationalSize") or "").strip().lower()
    implied = _LETTER.get(size.replace(" ", ""))
    has_render = any(m.get("view") in AI_VIEWS for m in media if m.get("isCurrent", True))
    if has_render and stored and implied and stored != implied:
        add("BODY_TYPE_UNRECORDED", "low",
            f"imageSettings.bodyType is '{stored}' but size {size.upper()} implies '{implied}'", "none")

    if measured:
        action = measured.get("action")
        if action == "bad":
            for reason in measured.get("reasons") or []:
                low = reason.lower()
                if ": canvas:" in low or ": frame:" in low:
                    add("CUTOUT_ZOOMED", "medium", reason, "rematte")
                elif ": backdrop:" in low:
                    add("CUTOUT_NOT_BACKDROP", "medium", reason, "rematte")
        elif action == "unknown":
            add("CUTOUT_UNMEASURED", "low", "; ".join(measured.get("reasons") or []), "none")

    if judged:
        code = judged.get("code")
        reasons = "; ".join(judged.get("reasons") or [])
        if code == "MODEL_GENDER_MISMATCH":
            add("RENDER_WRONG_GENDER", "high", reasons, "regen")
        elif code == "BODY_SIZE_MISMATCH":
            add("RENDER_BODY_MISMATCH", "high", reasons, "regen")
        elif code == "IMAGE_QUALITY" and judged.get("action") == "regen":
            add("RENDER_DEFECT", "high", reasons, "regen")
        elif code == "CATEGORY_IMAGE_MISMATCH":
            add("RENDER_CATEGORY_MISMATCH", "medium", reasons, "data")
        elif code == "VISION_UNAVAILABLE" or judged.get("unavailable"):
            add("RENDER_UNJUDGED", "low", reasons, "none")
        for soft in judged.get("soft") or []:
            if "one band off" in soft:
                add("RENDER_BUILD_ONE_BAND", "low", soft, "none")
    return issues


def worst(issues: list[dict[str, str]]) -> str:
    return min((i["severity"] for i in issues), key=lambda s: RANK.get(s, 3), default="none")


def fixes_for(issues: list[dict[str, str]]) -> list[str]:
    order = ["rematte", "regen", "reorder", "copy"]
    present = {i["fix"] for i in issues}
    return [f for f in order if f in present]


# --------------------------------------------------------------------------- #
# Gathering
# --------------------------------------------------------------------------- #

STAGE_SQL = """
SELECT p.id::text
  FROM "Product" p
  JOIN "Tenant" t ON t.id = p."tenantId"
 WHERE p."isDeleted" = false
   AND p."currentStage"::text = %(stage)s
   AND (%(tenant)s::text IS NULL OR t.name = %(tenant)s)
 ORDER BY p."createdAt" DESC
 LIMIT %(limit)s
"""


def _measure(loaded: dict[str, Any], pol: dict[str, Any]) -> dict[str, Any]:
    from app.imaging import cutouts
    from app.vnyx_client import to_snapshot

    snap = to_snapshot({**loaded["record"], "media": loaded["media"]},
                       catalog=loaded.get("catalog"), imagery_settings=loaded.get("imagery_settings"))
    return cutouts.judge(snap, pol).as_dict()


def _judge(loaded: dict[str, Any], pol: dict[str, Any]) -> dict[str, Any]:
    from app import readiness
    from app.imaging import quality_gate

    rec = loaded["record"]
    return quality_gate.judge(
        loaded["media"],
        gender=readiness.root_gender(rec.get("masterCategory"), pol) or rec.get("gender"),
        category=rec.get("category"), subcategory=rec.get("subCategory"), pol=pol,
        size=rec.get("size") or rec.get("internationalSize"),
        kids=readiness.is_kids(rec.get("masterCategory"), rec.get("mannequinType"), pol),
    ).as_dict()


def gather(dsn: str, *, stage: str, tenant: str | None, limit: int | None,
           start: datetime | None, end: datetime | None, product_ids: list[str] | None,
           measure: bool, judge: bool, workers: int, quiet: bool) -> list[dict[str, Any]]:
    from app import approval

    pol = policy()
    with product_audit.connect(dsn, read_only=True, statement_timeout_s=180) as conn, conn.cursor() as cur:
        if product_ids:
            ids = list(dict.fromkeys(product_ids))
        elif start and end:
            from scripts.audit_approved_window import APPROVED_SQL

            cur.execute(APPROVED_SQL, {"start": start, "end": end, "tenant": tenant})
            ids = [r[0] for r in cur.fetchall()]
            if limit:
                ids = ids[:limit]
        else:
            cur.execute(STAGE_SQL, {"stage": stage, "tenant": tenant, "limit": limit})
            ids = [r[0] for r in cur.fetchall()]
        if not quiet:
            print(f"{len(ids)} product(s) in scope"
                  + (f" ({stage})" if not (product_ids or start) else ""))
        if not ids:
            return []

        out: list[dict[str, Any]] = []
        contexts: dict[str, dict[str, Any]] = {}
        for i in range(0, len(ids), BATCH):
            chunk = ids[i:i + BATCH]
            cur.execute('SELECT DISTINCT "tenantId"::text FROM "Product" WHERE id = ANY(%s::uuid[])', (chunk,))
            for (tid,) in cur.fetchall():
                if tid not in contexts:
                    contexts[tid] = product_audit.load_tenant_context(cur, tid)
            loaded_list = product_audit.load_batch(cur, chunk, contexts)

            measured: dict[str, dict[str, Any]] = {}
            judged: dict[str, dict[str, Any]] = {}
            if measure or judge:
                def one(l: dict[str, Any]) -> tuple[str, dict[str, Any] | None, dict[str, Any] | None]:
                    pid = l["record"]["id"]
                    m = _measure(l, pol) if measure else None
                    j = _judge(l, pol) if judge else None
                    return pid, m, j
                with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
                    for pid, m, j in pool.map(one, loaded_list):
                        if m is not None:
                            measured[pid] = m
                        if j is not None:
                            judged[pid] = j

            for l in loaded_list:
                record, media = l["record"], l["media"]
                pid = record["id"]
                gate = approval.run_gate({**record, "media": media}, catalog=l["catalog"],
                                         imagery_settings=l["imagery_settings"], llm=None)
                findings = [*gate["blocking"], *gate["advisory"]]
                issues = assess_readiness(record, media, l.get("chart"), findings,
                                          measured=measured.get(pid), judged=judged.get(pid))
                live = [m for m in media if m.get("isCurrent", True) and m.get("mediaType", "IMAGE") == "IMAGE"]
                out.append({
                    "pid": pid, "sku": record.get("sku"), "title": record.get("title"),
                    "tenant": record.get("tenantName"), "stage": record.get("currentStage"),
                    "master": record.get("masterCategory"), "category": record.get("category"),
                    "subcategory": record.get("subCategory"), "size": record.get("size"),
                    "renders": sum(1 for v in AI_VIEWS if v in {m.get("view") for m in live}),
                    "cutouts": sum(1 for m in live if m.get("view") in ("FRONT", "BACK")
                                   and m.get("processing") != "RAW"),
                    "labels": record.get("careLabelCount") or 0,
                    "issues": issues, "worst": worst(issues), "fixes": fixes_for(issues),
                    "measured": measured.get(pid), "judged": judged.get(pid),
                    "edit_url": record.get("editUrl"),
                })
            if not quiet:
                print(f"  {min(i + BATCH, len(ids))}/{len(ids)} examined")
    return out


# --------------------------------------------------------------------------- #
# Outputs
# --------------------------------------------------------------------------- #

FILL = {"high": "FDECEA", "medium": "FFF8E1", "low": "EEF2FB", "none": "E8F5E9"}
PRODUCT_COLUMNS = [("Worst", 9), ("Fixes", 22), ("Issue codes", 46), ("SKU", 14), ("Title", 40),
                   ("Tenant", 14), ("Stage", 10), ("Master", 8), ("Category", 18), ("Subcategory", 18),
                   ("Size", 6), ("Renders", 8), ("Cut-outs", 8), ("Labels", 7), ("Product id", 38),
                   ("Edit URL", 60)]
ISSUE_COLUMNS = [("Severity", 9), ("Code", 30), ("Fix", 9), ("SKU", 14), ("Tenant", 14),
                 ("Detail", 100), ("Product id", 38), ("Edit URL", 60)]


def write_workbook(rows: list[dict[str, Any]], out: Path, *, scope: str, method: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font = Font(bold=True, color="FFFFFF")

    def header(ws, columns, row=1):
        for i, (name, width) in enumerate(columns, start=1):
            c = ws.cell(row=row, column=i, value=name)
            c.fill, c.font = hdr_fill, hdr_font
            ws.column_dimensions[get_column_letter(i)].width = width

    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Readiness audit").font = Font(bold=True, size=14)
    counts = Counter(i["code"] for r in rows for i in r["issues"])
    affected = Counter(i["code"] for r in rows for i in {x["code"]: x for x in r["issues"]}.values())
    sev_of = {i["code"]: i["severity"] for r in rows for i in r["issues"]}
    fix_of = {i["code"]: i["fix"] for r in rows for i in r["issues"]}
    by_fix = Counter(f for r in rows for f in r["fixes"])
    meta = [("Scope", scope), ("Products", len(rows)),
            ("With at least one issue", sum(1 for r in rows if r["issues"])),
            ("Worst = high", sum(1 for r in rows if r["worst"] == "high")),
            ("Worst = medium", sum(1 for r in rows if r["worst"] == "medium")),
            ("To re-matte", by_fix.get("rematte", 0)), ("To re-render", by_fix.get("regen", 0)),
            ("To re-order", by_fix.get("reorder", 0)), ("To re-write", by_fix.get("copy", 0)),
            ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
            ("Method", method)]
    for n, (k, v) in enumerate(meta, start=3):
        ws.cell(row=n, column=1, value=k).font = Font(bold=True)
        ws.cell(row=n, column=2, value=v)
    r0 = 3 + len(meta) + 1
    header(ws, [("Issue code", 34), ("Severity", 10), ("Fix", 9), ("Occurrences", 13),
                ("Products affected", 18), ("Meaning", 90)], row=r0)
    for n, (code, cnt) in enumerate(sorted(counts.items(), key=lambda kv: (RANK[sev_of[kv[0]]], -kv[1])),
                                    start=r0 + 1):
        values = [code, sev_of[code], fix_of[code], cnt, affected[code],
                  MEANING.get(code.split(":")[0], MEANING["RULE"])]
        for col, v in enumerate(values, start=1):
            ws.cell(row=n, column=col, value=v).fill = PatternFill("solid", fgColor=FILL[sev_of[code]])

    ws = wb.create_sheet("Products")
    header(ws, PRODUCT_COLUMNS)
    ws.freeze_panes = "A2"
    for n, r in enumerate(sorted(rows, key=lambda r: (RANK.get(r["worst"], 3), -len(r["issues"]), str(r["sku"]))),
                          start=2):
        values = [r["worst"], ", ".join(r["fixes"]) or "—", ", ".join(sorted({i["code"] for i in r["issues"]})),
                  r["sku"], (r["title"] or "")[:80], r["tenant"], r["stage"], r["master"], r["category"],
                  r["subcategory"], r["size"], f'{r["renders"]}/5', r["cutouts"], r["labels"], r["pid"],
                  r["edit_url"] or ""]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=n, column=col, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=col in (3, 5))
        ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=FILL.get(r["worst"], FILL["none"]))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PRODUCT_COLUMNS))}{max(ws.max_row, 2)}"

    ws = wb.create_sheet("Issues")
    header(ws, ISSUE_COLUMNS)
    ws.freeze_panes = "A2"
    flat = [(r, i) for r in rows for i in r["issues"]]
    flat.sort(key=lambda ri: (RANK[ri[1]["severity"]], ri[1]["code"], str(ri[0]["sku"])))
    for n, (r, i) in enumerate(flat, start=2):
        values = [i["severity"], i["code"], i["fix"], r["sku"], r["tenant"], i["detail"], r["pid"], r["edit_url"] or ""]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=n, column=col, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=col == 6)
        ws.cell(row=n, column=1).fill = PatternFill("solid", fgColor=FILL[i["severity"]])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ISSUE_COLUMNS))}{max(ws.max_row, 2)}"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def flagged(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Product ids per fix — the backlog's input (decision 7)."""
    lists: dict[str, list[str]] = {"rematte": [], "regen": [], "reorder": [], "copy": []}
    for r in rows:
        for f in r["fixes"]:
            lists[f].append(r["pid"])
    return {
        "generated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "products": len(rows),
        "counts": {k: len(v) for k, v in lists.items()},
        "commands": FIX_COMMANDS,
        **lists,
    }


def _redact(dsn: str) -> str:
    try:
        head, tail = dsn.split("@", 1)
        return head.split("//", 1)[0] + "//***@" + tail
    except ValueError:
        return dsn


def _day(text: str) -> datetime:
    return datetime.combine(date.fromisoformat(text), datetime.min.time())


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", "--dsn", dest="db", help="database URL; else DATABASE_URL")
    ap.add_argument("--stage", default="APPROVED", help="products in this stage (default APPROVED)")
    ap.add_argument("--tenant")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--from", dest="start", metavar="YYYY-MM-DD",
                    help="with --to: products that entered APPROVED in the window instead of --stage")
    ap.add_argument("--to", dest="end", metavar="YYYY-MM-DD", help="inclusive")
    ap.add_argument("--product", "--products", dest="products", help="uuid[,uuid] — these, whatever their stage")
    ap.add_argument("--measure-images", action="store_true",
                    help="download and measure the cut-outs (canvas, backdrop); no model call")
    ap.add_argument("--judge-images", action="store_true",
                    help="put the lead render through the image gate (one vision call per product, cached)")
    ap.add_argument("--workers", type=int, default=4, help="parallel image reads / vision calls")
    ap.add_argument("--out", help="workbook path (default reports/generated/readiness-<stamp>.xlsx)")
    ap.add_argument("--flagged", help="write the per-fix product-id lists here (JSON)")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    # The rule messages carry × and ΔE; a cp1252 console must not end the run.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except Exception:  # noqa: BLE001
        pass

    from app.config import settings

    dsn = args.db or settings().database_url
    if not dsn:
        print("Pass --db or set DATABASE_URL.", file=sys.stderr)
        return 2
    if bool(args.start) != bool(args.end):
        print("--from and --to go together.", file=sys.stderr)
        return 2
    start = _day(args.start) if args.start else None
    end = _day(args.end) + timedelta(days=1) if args.end else None
    product_ids = [p.strip() for p in (args.products or "").split(",") if p.strip()] or None

    rows = gather(dsn, stage=args.stage.upper(), tenant=args.tenant, limit=args.limit,
                  start=start, end=end, product_ids=product_ids,
                  measure=args.measure_images, judge=args.judge_images,
                  workers=args.workers, quiet=args.quiet)
    if not rows:
        print("Nothing in scope.")
        return 0

    scope = (f"{len(product_ids)} named product(s)" if product_ids
             else f"entered APPROVED {args.start}..{args.end}" if start
             else f"stage {args.stage.upper()}" + (f", tenant {args.tenant}" if args.tenant else "")
             + (f", newest {args.limit}" if args.limit else ""))
    method = "rules over the database" + (", cut-outs measured" if args.measure_images else "") \
        + (", lead renders judged by the image gate" if args.judge_images else "") + "; nothing written"
    stamp = datetime.now().strftime("%Y-%m-%d-%H%M%S")
    out = Path(args.out) if args.out else Path("reports/generated") / f"readiness-{stamp}.xlsx"
    write_workbook(rows, out, scope=scope, method=method)

    fl = flagged(rows)
    if args.flagged:
        Path(args.flagged).parent.mkdir(parents=True, exist_ok=True)
        Path(args.flagged).write_text(json.dumps(fl, indent=1), encoding="utf-8")

    counts = Counter(i["code"] for r in rows for i in r["issues"])
    print(f"\n{len(rows)} product(s) · {sum(1 for r in rows if r['issues'])} with an issue · "
          f"worst high {sum(1 for r in rows if r['worst'] == 'high')} · database {_redact(dsn)}")
    for code, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {n:5}  {code}")
    print("\nbacklog (decision 7 — only these):")
    for fix in ("rematte", "regen", "reorder", "copy"):
        print(f"  {fix:8} {fl['counts'][fix]:5}  {FIX_COMMANDS[fix]}")
    print(f"\nworkbook: {out}" + (f"\nflagged : {args.flagged}" if args.flagged else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
