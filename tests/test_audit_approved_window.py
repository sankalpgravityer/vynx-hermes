"""scripts/audit_approved_window.py — the decision table, cold.

`assess()` takes plain dicts, so every issue code is exercised here with no
database: who approved, what the run record says the gate and the photo audit
decided, what is on file, and what the channel tables say.
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts import audit_approved_window as aw  # noqa: E402

MEDIA_OK = (
    [{"url": f"https://r2/{v}.jpg", "view": v, "processing": "GENERATED", "mediaType": "IMAGE"} for v in aw.AI_VIEWS]
    + [{"url": "https://r2/front-cut.png", "view": "FRONT", "processing": "BG_REMOVED", "mediaType": "IMAGE"},
       {"url": "https://r2/back-cut.png", "view": "BACK", "processing": "BG_REMOVED", "mediaType": "IMAGE"},
       {"url": "https://r2/label.jpg", "view": "LABEL", "processing": "RAW", "mediaType": "IMAGE"}]
)
RECORD_OK = {"sku": "BOA-1", "tenantId": "t1", "generationStatus": "COMPLETE", "isRegenerating": False,
             "careLabelCount": 1, "images": ["https://r2/AI_FRONT.jpg"],
             "verificationStatus": "VERIFIED", "verificationOutcome": "APPROVED"}
AGENT_OK = {"status": "VERIFIED", "outcome": "APPROVED", "approved": True,
            "deltas": {"gate": {"action": "ok", "code": None, "reasons": []},
                       "photos": {"action": "ok", "code": None, "reasons": []}}}
LISTING_OK = [{"code": "BOAS", "status": "PUBLISHED", "lastSyncStatus": "SUCCESS", "externalListingId": "1"}]
APPROVED = {"at": datetime(2026, 9, 15, 14, 22), "by": None, "from_stage": "REVIEW"}


def _assess(**over):
    args = dict(record=RECORD_OK, media=MEDIA_OK, approved=APPROVED, agent=AGENT_OK, blocking=[],
                listings=LISTING_OK, logs={"publication-verify": {"status": "SUCCESS"}},
                has_account=True, shopify_id="10", twin_gains=[])
    args.update(over)
    return aw.assess(**args)


def codes(issues):
    return sorted(i["code"] for i in issues)


def test_a_product_the_agent_approved_with_every_check_passing_has_no_issue():
    assert _assess() == []


def test_a_manual_approval_never_judged_by_the_agent():
    issues = _assess(agent=None, approved={**APPROVED, "by": "Sankalp"},
                     record={**RECORD_OK, "verificationStatus": "NOT_STARTED", "verificationOutcome": None})
    assert codes(issues) == ["MANUAL_APPROVAL", "NO_GATE_RECORD", "NO_PHOTO_AUDIT"]
    manual = next(i for i in issues if i["code"] == "MANUAL_APPROVAL")
    assert "Sankalp" in manual["detail"] and "NOT_STARTED" in manual["detail"]
    assert aw.worst(issues) == "medium"


def test_a_human_overriding_a_hold_is_named():
    held = {**AGENT_OK, "status": "HELD_FOR_HUMAN", "outcome": "PRICE.003", "approved": False}
    issues = _assess(agent=held, approved={**APPROVED, "by": "Sankalp"})
    assert "AGENT_HELD_BUT_APPROVED" in codes(issues) and "MANUAL_APPROVAL" in codes(issues)


def test_gate_and_photo_audit_refusals_that_were_approved_anyway_are_high():
    agent = {**AGENT_OK, "deltas": {
        "gate": {"action": "regen", "code": "MODEL_GENDER_MISMATCH", "reasons": ["women on a men's product"]},
        "photos": {"action": "review", "code": "GRADE_SUSPECT", "reasons": ["major wear vs none"]}}}
    issues = _assess(agent=agent)
    assert codes(issues) == ["GATE_REFUSED_BUT_APPROVED", "PHOTO_AUDIT_HELD_BUT_APPROVED"]
    assert all(i["severity"] == "high" for i in issues)
    assert aw.worst(issues) == "high"


def test_blocking_rules_become_rule_issues_with_their_severity():
    blocking = [{"rule_id": "PRICE.003", "severity": "high", "message": "13.99 is only 22% of retail"},
                {"rule_id": "DATA.003", "severity": "medium", "message": "Inventory is 2"}]
    issues = _assess(blocking=blocking)
    by = {i["code"]: i for i in issues}
    assert by["RULE:PRICE.003"]["severity"] == "high" and "22%" in by["RULE:PRICE.003"]["detail"]
    assert by["RULE:DATA.003"]["severity"] == "medium"


def test_advisory_findings_are_itemised_at_their_own_severity_and_conf_is_folded():
    advisory = [{"rule_id": "DATA.003", "severity": "medium", "message": "Inventory is 2"},
                {"rule_id": "TEXT.007", "severity": "medium", "message": "copy says men"},
                {"rule_id": "TEXT.005", "severity": "low", "message": "material is a sentence"},
                {"rule_id": "CONF.001", "severity": "low", "message": "fit 65%", "fields": ["fit"]},
                {"rule_id": "CONF.001", "severity": "low", "message": "waist 50%", "fields": ["waist"]}]
    issues = _assess(advisory=advisory)
    by = {i["code"]: i for i in issues}
    assert by["RULE:DATA.003"]["severity"] == "medium" and by["RULE:TEXT.005"]["severity"] == "low"
    assert sum(1 for i in issues if i["code"] == "RULE:CONF.001") == 1
    assert "2 field(s)" in by["RULE:CONF.001"]["detail"] and "fit, waist" in by["RULE:CONF.001"]["detail"]
    assert aw.worst(issues) == "medium"


def test_judged_now_a_refusing_gate_and_a_holding_photo_audit_are_high():
    judged = {"gate": {"action": "regen", "code": "MODEL_GENDER_MISMATCH", "reasons": ["women on men's"],
                       "unavailable": False},
              "photos": {"action": "review", "code": "GRADE_SUSPECT", "reasons": ["major wear vs none"]}}
    issues = _assess(agent={**AGENT_OK, "deltas": {}}, judged=judged)
    assert codes(issues) == ["GATE_WOULD_REFUSE", "PHOTO_AUDIT_WOULD_HOLD"]
    assert all(i["severity"] == "high" for i in issues)
    assert "judged now" in issues[0]["detail"]


def test_judged_now_a_defective_render_is_named_with_the_view_to_re_render():
    """MID-000253 in the approved-window sheet: 'judged now: ok (IMAGE DEFECT —
    AI_FRONT_34 render: AI artifacts on legs)' was a soft flag in a low row.
    The same answer is now a high RENDER_DEFECT that says what to re-render."""
    judged = {"gate": {"action": "ok", "code": None, "reasons": [], "unavailable": False},
              "photos": {"action": "regen", "code": "RENDER_DEFECT",
                         "reasons": ["RENDER DEFECT — AI_FRONT_34 render: AI artifact on legs"],
                         "bad_views": ["AI_FRONT_34"]}}
    issues = _assess(agent={**AGENT_OK, "deltas": {}}, judged=judged)
    assert codes(issues) == ["RENDER_DEFECT"]
    assert issues[0]["severity"] == "high"
    assert "re-render AI_FRONT_34" in issues[0]["detail"] and "same model" in issues[0]["detail"]
    assert "RENDER_DEFECT" in aw.MEANING
    # Recorded by the agent as a refusal and approved anyway: the same high row as a hold.
    agent = {**AGENT_OK, "deltas": {"gate": {"action": "ok"}, "photos": judged["photos"]}}
    assert codes(_assess(agent=agent)) == ["PHOTO_AUDIT_HELD_BUT_APPROVED"]


def test_judged_now_a_cut_out_missing_its_collar_is_a_medium_row_naming_the_re_cut():
    """MID-000569: the neckband cut away on both FRONT cut-outs — a flag while
    the cut-out policy is soft, never buried among the soft flags."""
    judged = {"gate": {"action": "ok", "code": None, "reasons": [], "unavailable": False},
              "photos": {"action": "ok", "code": None, "reasons": [],
                         "soft": ["CUTOUT DEFECT — FRONT cut-out: collar cut away by background removal"],
                         "bad_cutouts": ["FRONT"]}}
    issues = _assess(agent={**AGENT_OK, "deltas": {}}, judged=judged)
    assert codes(issues) == ["CUTOUT_DEFECT"]
    assert issues[0]["severity"] == "medium" and "re-cut FRONT" in issues[0]["detail"]
    assert "CUTOUT_DEFECT" in aw.MEANING


def test_judged_now_a_clean_pass_leaves_only_soft_flags_and_no_never_judged_codes():
    judged = {"gate": {"action": "ok", "code": None, "reasons": [], "unavailable": False},
              "photos": {"action": "ok", "code": None, "reasons": [], "soft": ["IMAGE DEFECT — AI_BACK render: clutter"]}}
    issues = _assess(agent=None, approved={**APPROVED, "by": "Sankalp"}, judged=judged)
    assert codes(issues) == ["MANUAL_APPROVAL", "PHOTO_AUDIT_SOFT_FLAGS"]
    assert "clutter" in next(i for i in issues if i["code"] == "PHOTO_AUDIT_SOFT_FLAGS")["detail"]


def test_judged_now_but_the_provider_was_down_is_its_own_finding():
    """An outage is a fault in the RUN, and says nothing about the product.

    Two things it must not be confused with, which is the whole point of giving
    it a code of its own:

      * a refusal. `unavailable` is not a verdict, so GATE_WOULD_REFUSE stays out.
      * `NO_GATE_RECORD` / `NO_PHOTO_AUDIT`, which mean "nobody ever judged
        this" — history a reader discounts. A sheet where the provider was down
        for every product then reads as an ordinary backlog of unjudged ones,
        and that is exactly what happened on 18 Sep 2026 when `google-genai` was
        missing from the environment: 393 products, no refusals, nothing in the
        summary to say the pictures had not actually been looked at.

    HIGH, because an unchecked picture on an approved product is the thing this
    report exists to find.
    """
    judged = {"gate": {"action": "review", "code": "VISION_UNAVAILABLE", "reasons": ["429"], "unavailable": True},
              "photos": {"action": "skipped", "code": None, "reasons": ["could not run — 429"]}}
    issues = _assess(agent={**AGENT_OK, "deltas": {}}, judged=judged)
    assert codes(issues) == ["PICTURES_COULD_NOT_BE_JUDGED", "PICTURES_COULD_NOT_BE_JUDGED"]
    assert all(i["severity"] == "high" for i in issues)
    assert {i["detail"].split(" could not run")[0] for i in issues} == {"image gate", "photo audit"}
    assert all("429" in i["detail"] for i in issues)
    assert "PICTURES_COULD_NOT_BE_JUDGED" in aw.MEANING


def test_pictures_on_file_incomplete_renders_no_label_unmatted_and_lead():
    media = [m for m in MEDIA_OK if m["view"] not in ("AI_CLOSEUP", "LABEL", "BACK")]
    media += [{"url": "https://r2/back-raw.jpg", "view": "BACK", "processing": "RAW", "mediaType": "IMAGE"},
              {"url": "https://r2/front-raw.jpg", "view": "FRONT", "processing": "RAW", "mediaType": "IMAGE"}]
    record = {**RECORD_OK, "careLabelCount": 0, "images": ["https://r2/front-raw.jpg"],
              "generationStatus": "GENERATING"}
    issues = _assess(media=media, record=record)
    assert {"RENDERS_INCOMPLETE", "NO_CARE_LABEL", "UNMATTED_VIEW", "RAW_LEAD",
            "GENERATION_NOT_COMPLETE"} <= set(codes(issues))
    renders = next(i for i in issues if i["code"] == "RENDERS_INCOMPLETE")
    assert "4/5" in renders["detail"] and "AI_CLOSEUP" in renders["detail"]
    assert "BACK" in next(i for i in issues if i["code"] == "UNMATTED_VIEW")["detail"]


def test_a_label_lead_is_named():
    record = {**RECORD_OK, "images": ["https://r2/label.jpg"]}
    assert "LABEL_LEAD" in codes(_assess(record=record))


def test_twin_blanks_the_parent_holds():
    issues = _assess(twin_gains=["size", "brand"])
    assert codes(issues) == ["TWIN_BLANK"] and "size, brand" in issues[0]["detail"]


def test_channel_checks_only_when_the_tenant_has_an_account():
    assert _assess(listings=[], logs={}, has_account=False) == []
    issues = _assess(listings=[], logs={}, has_account=True)
    assert codes(issues) == ["NO_LISTING"]


def test_listing_states_failed_not_published_unverified_and_no_shopify_id():
    listings = [{"code": "BOAS", "status": "IDLE", "lastSyncStatus": "FAILED", "externalListingId": None,
                 "lastSyncError": "publishablePublish: Publication does not exist"}]
    logs = {"publication-verify": {"status": "FAILED", "errorMessage": "could not verify publication"}}
    issues = _assess(listings=listings, logs=logs, shopify_id=None)
    assert codes(issues) == ["LISTING_NOT_PUBLISHED", "LISTING_SYNC_FAILED", "PUBLICATION_UNVERIFIED"]
    assert all(i["severity"] == "high" for i in issues)
    published_no_id = [{"code": "BOAS", "status": "PUBLISHED", "lastSyncStatus": "SUCCESS", "externalListingId": None}]
    assert codes(_assess(listings=published_no_id, shopify_id=None)) == ["NO_SHOPIFY_ID"]


def test_workbook_is_written_with_three_sheets(tmp_path):
    from openpyxl import load_workbook

    rows = [{
        "pid": "p1", "sku": "BOA-1", "title": "Jeans", "tenant": "BOAS",
        "approved_at": datetime(2026, 9, 15, 14, 22), "approved_by": "Sankalp", "stage_now": "APPROVED",
        "verification": "NOT_STARTED", "verification_outcome": None, "agent": None, "gate": None, "photos": None,
        "renders": 5, "care_labels": 0, "listings": LISTING_OK, "shopify_id": "10",
        "publication_verify": "FAILED", "blocking": ["PRICE.003"], "advisory": 3,
        "issues": _assess(agent=None, record={**RECORD_OK, "careLabelCount": 0},
                          logs={"publication-verify": {"status": "FAILED", "errorMessage": "x"}}),
        "worst": "high", "edit_url": "https://try.vnyx.ai/product/p1/edit",
    }]
    out = aw.write_workbook(rows, tmp_path / "a.xlsx", dsn="postgresql://u:pw@host/db",
                            start=datetime(2026, 9, 10), end=datetime(2026, 9, 16), tenant=None)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Products", "Issues"]
    summary = [c.value for c in wb["Summary"]["B"]]
    assert "postgresql://***@host/db" in summary and "pw" not in str(summary)
    products = wb["Products"]
    assert products.cell(row=2, column=4).value == "BOA-1" and products.cell(row=2, column=1).value == "high"
    issues = wb["Issues"]
    assert issues.max_row - 1 == len(rows[0]["issues"])
    assert issues.cell(row=2, column=1).value == "high"


def _fake_gather(seen):
    def fake_gather(dsn, start, end, tenant, limit, quiet=False, judge_images=False, workers=1,
                    product_ids=None, rejudge=False):
        seen.update(start=start, end=end, tenant=tenant, judge_images=judge_images, workers=workers,
                    product_ids=product_ids, rejudge=rejudge)
        return []
    return fake_gather


def test_window_parsing_is_inclusive(monkeypatch, tmp_path):
    seen = {}
    monkeypatch.setattr(aw, "gather", _fake_gather(seen))
    assert aw.main(["--db", "postgresql://x", "--from", "2026-09-10", "--to", "2026-09-15",
                    "--tenant", "BOAS", "--out", str(tmp_path / "w.xlsx"), "--quiet"]) == 0
    assert seen["start"] == datetime(2026, 9, 10) and seen["end"] == datetime(2026, 9, 16)
    assert seen["tenant"] == "BOAS" and seen["judge_images"] is False and seen["rejudge"] is False
    assert (tmp_path / "w.xlsx").exists()


PID = "0f1e2d3c-4b5a-4978-8877-665544332211"
PID2 = "11111111-2222-4333-8444-555555555555"


def test_product_mode_takes_ids_and_rejudge_implies_judging(monkeypatch, tmp_path):
    seen = {}
    monkeypatch.setattr(aw, "gather", _fake_gather(seen))
    # --from/--to left on the command line are ignored, not an error: the user
    # appends --product to the window command they already have.
    assert aw.main(["--db", "postgresql://x", "--from", "2026-09-10", "--to", "2026-09-15",
                    "--product", f"{PID},{PID2}", "--rejudge",
                    "--out", str(tmp_path / "p.xlsx"), "--quiet"]) == 0
    assert seen["product_ids"] == [PID, PID2]
    assert seen["start"] is None and seen["end"] is None
    assert seen["judge_images"] is True and seen["rejudge"] is True
    assert (tmp_path / "p.xlsx").exists()


def test_rejudge_needs_product_and_product_needs_uuids(monkeypatch, tmp_path):
    import pytest

    monkeypatch.setattr(aw, "gather", _fake_gather({}))
    with pytest.raises(SystemExit) as exc:
        aw.main(["--db", "postgresql://x", "--from", "2026-09-10", "--to", "2026-09-15", "--rejudge"])
    assert "--product" in str(exc.value)
    with pytest.raises(SystemExit) as exc:
        aw.main(["--db", "postgresql://x", "--product", "BOA-1"])
    assert "BOA-1" in str(exc.value)


class _Verdict:
    def __init__(self, **kw):
        self.kw = kw

    def as_dict(self):
        return dict(self.kw)


def test_judge_now_judges_only_what_is_missing_unless_forced(monkeypatch):
    from app.imaging import photo_audit, quality_gate

    calls = []
    monkeypatch.setattr(quality_gate, "judge", lambda media, **kw: calls.append("gate") or _Verdict(action="ok"))
    monkeypatch.setattr(photo_audit, "judge", lambda media, **kw: calls.append("photos") or _Verdict(action="ok"))
    record = {"gender": "men", "category": "Tops", "subCategory": "T-shirt", "gradeSeverity": 1}

    out = aw.judge_now(record, MEDIA_OK, AGENT_OK, pol={})
    assert out == {} and calls == []

    only_gate = {**AGENT_OK, "deltas": {"gate": AGENT_OK["deltas"]["gate"]}}
    out = aw.judge_now(record, MEDIA_OK, only_gate, pol={})
    assert set(out) == {"photos"} and calls == ["photos"]

    calls.clear()
    out = aw.judge_now(record, MEDIA_OK, AGENT_OK, pol={}, force=True)
    assert set(out) == {"gate", "photos"} and sorted(calls) == ["gate", "photos"]


def test_fresh_verdicts_replace_recorded_ones_in_what_assess_reads():
    recorded_refusal = {**AGENT_OK, "deltas": {
        "gate": {"action": "regen", "code": "MODEL_GENDER_MISMATCH", "reasons": ["women on men's"]},
        "photos": {"action": "ok", "code": None, "reasons": []}}}
    # Nothing judged now: the record is read as it is.
    assert aw.fresh_over_recorded(recorded_refusal, {}) is recorded_refusal
    assert aw.fresh_over_recorded(None, {"gate": {"action": "ok"}}) is None

    judged = {"gate": {"action": "ok", "code": None, "reasons": [], "unavailable": False}}
    as_read = aw.fresh_over_recorded(recorded_refusal, judged)
    assert "gate" not in as_read["deltas"] and "photos" in as_read["deltas"]
    assert as_read["approved"] is True and as_read["status"] == "VERIFIED"
    # The gate passes today, so the old refusal is not reported; nothing else changes.
    assert _assess(agent=as_read, judged=judged) == []

    judged = {"gate": {"action": "regen", "code": "IMAGE_QUALITY", "reasons": ["no model"], "unavailable": False},
              "photos": {"action": "review", "code": "IMAGE_DEFECT", "reasons": ["clutter"]}}
    issues = _assess(agent=aw.fresh_over_recorded(AGENT_OK, judged), judged=judged)
    assert codes(issues) == ["GATE_WOULD_REFUSE", "PHOTO_AUDIT_WOULD_HOLD"]
    assert all("judged now" in i["detail"] for i in issues)


def test_verdict_line_reads_like_the_step_note():
    assert aw.verdict_line(None) == "not judged"
    assert aw.verdict_line({"action": "ok", "cached": True}) == "passed (cached)"
    assert aw.verdict_line({"action": "regen", "code": "MODEL_GENDER_MISMATCH", "reasons": ["women on men's"]}) \
        == "REFUSED MODEL_GENDER_MISMATCH — women on men's"
    assert aw.verdict_line({"action": "review", "code": "GRADE_SUSPECT", "reasons": ["major wear vs none"],
                            "soft": ["AI_BACK: clutter"]}) \
        == "HELD GRADE_SUSPECT — major wear vs none (soft: AI_BACK: clutter)"
    assert aw.verdict_line({"action": "review", "code": "VISION_UNAVAILABLE", "reasons": ["429"],
                            "unavailable": True}) == "could not run VISION_UNAVAILABLE — 429"
